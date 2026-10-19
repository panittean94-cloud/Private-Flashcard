# -*- coding: utf-8 -*-
"""
clean_and_assemble.py
Sanitizes strings, ensures pure text questions, fixes illegal control chars,
maps pure diagrams, and generates:
1. data/flashcards.json
2. data/flashcards.js
3. FunE_Master_Exam_Flashcards_Ch6_10_Master.xlsx
"""

import os
import re
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import create_bilingual_part1
import create_bilingual_part2
import create_bilingual_part3

def strip_embedded_html(text):
    if not isinstance(text, str):
        return text
    # Remove div/img wrappers
    text = re.sub(r'<div class=[\'\"].*?</div></div>', '', text, flags=re.DOTALL)
    text = re.sub(r'<div class=[\'\"].*?</div>', '', text, flags=re.DOTALL)
    # Fix vertical tabs or weird control chars
    text = text.replace('\x0b', '\\vec')
    # Strip trailing/leading spaces
    return text.strip()

all_cards = create_bilingual_part2.full_cards + create_bilingual_part3.cards_p3
print(f"Loaded total {len(all_cards)} raw cards")

diagram_map = {
    "CH06-04": ("images/diagrams/hw6_fig1_yy_balanced.png", "Homework 6 Figure 1: Balanced Y-Y System", "การบ้านบทที่ 6 รูปที่ 1: ระบบ 3 เฟสแบบ Y-Y สมดุล"),
    "CH06-05": ("images/diagrams/hw6_fig2_yy_unbalanced.png", "Homework 6 Figure 2: Unbalanced Y-Y 4-Wire with Neutral Line", "การบ้านบทที่ 6 รูปที่ 2: ระบบ 3 เฟสแบบ Y-Y ไม่สมดุล มีสายนิวทรัล"),
    "CH06-06": ("images/diagrams/hw6_fig3_yy_neutral_z.png", "Homework 6 Figure 3: Unbalanced Y-Y 3-Wire System", "การบ้านบทที่ 6 รูปที่ 3: ระบบ 3 เฟสแบบ Y-Y ไม่สมดุล ไม่มีสายนิวทรัล"),
    "CH07-01": ("images/diagrams/exam_p28_prob41_right_hand.png", "Exam Problem 4.1: Right-Hand Grip Rule", "ข้อสอบจริง ข้อ 4.1: กฎมือขวาของสนามแม่เหล็กรอบเส้นลวด"),
    "CH07-02": ("images/diagrams/exam_p28_prob42_motor_force.png", "Exam Problem 4.2: DC Motor Magnetic Field & Force (Lorentz)", "ข้อสอบจริง ข้อ 4.2: สนามแม่เหล็กและแรงลอเรนซ์ในมอเตอร์ไฟฟ้ากระแสตรง"),
    "CH07-03": ("images/diagrams/exam_p28_prob43_motor_poles.png", "Exam Problem 4.3: Motor Magnetic Poles (N & S)", "ข้อสอบจริง ข้อ 4.3: ขั้วแม่เหล็ก N และ S ของมอเตอร์ไฟฟ้า"),
    "CH07-04": ("images/diagrams/hw7_p4_shunt_circuit.png", "Homework 7 Problem 4: DC Shunt Motor Equivalent Circuit", "การบ้านบทที่ 7 ข้อ 4: วงจรสมมูลมอเตอร์ไฟฟ้ากระแสตรงแบบชันต์ (DC Shunt Motor)"),
    "CH07-05": ("images/diagrams/hw7_p5_series_circuit.png", "Homework 7 Problem 5: DC Series Motor Equivalent Circuit", "การบ้านบทที่ 7 ข้อ 5: วงจรสมมูลมอเตอร์ไฟฟ้ากระแสตรงแบบอนุกรม (DC Series Motor)"),
    "CH07-06": ("images/diagrams/exam_p27_prob44_induction_motor.png", "Exam Problem 4.4: 3-Phase Induction Motor (Squirrel Cage)", "ข้อสอบจริง ข้อ 4.4: โครงสร้างมอเตอร์เหนี่ยวนำ 3 เฟส (โรเตอร์กรงกระรอก)"),
    "CH07-07": ("images/diagrams/exam_p26_prob45_generator.png", "Exam Problem 4.5: DC Generator Principle & Split-Ring Commutator", "ข้อสอบจริง ข้อ 4.5: เครื่องกำเนิดไฟฟ้ากระแสตรงและคอมมิวเทเตอร์"),
    "CH08-01": ("images/diagrams/exam_p32_transformer_core.png", "Exam Problem 1.2: Transformer Core & Dual Windings", "ข้อสอบจริง ข้อ 1.2: แกนเหล็กและขดลวดปฐมภูมิ-ทุติยภูมิของหม้อแปลง"),
    "CH08-02": ("images/diagrams/exam_p31_prob16_transformer_calc.png", "Exam Problem 1.6: 25 kVA Transformer Full Calculation", "ข้อสอบจริง ข้อ 1.6: การคำนวณหม้อแปลง 25 kVA ครบวงจร"),
    "CH08-08": ("images/diagrams/transformer_audio_matching.png", "Audio Transformer Maximum Power Transfer Impedance Matching", "หม้อแปลงขยายเสียงสำหรับปรับอิมพีแดนซ์ให้เกิดการส่งถ่ายกำลังสูงสุด"),
    "CH09-01": ("images/diagrams/power_system_stages.png", "Electric Power System Stages (Generation to Utilization)", "4 ขั้นตอนหลักของระบบไฟฟ้ากำลัง (ผลิต ส่ง จำหน่าย ใช้)"),
    "CH09-04": ("images/diagrams/exam_p25_prob51_52_power_plants.png", "Exam Problem 5.1 & 5.2: Power Plants in Thailand", "ข้อสอบจริง ข้อ 5.1 & 5.2: โรงไฟฟ้าและการใช้เชื้อเพลิงในไทย"),
    "CH09-05": ("images/diagrams/transmission_towers.png", "High-Voltage Transmission Towers (Suspension, Tension, Transposition)", "เสาสายส่งไฟฟ้าแรงสูง (เสาแขวน, เสารั้ง, เสาสลับเฟส)"),
    "CH09-06": ("images/diagrams/exam_p25_prob53_voltages.png", "Exam Problem 5.3 & 5.4: Grid Voltage Classifications", "ข้อสอบจริง ข้อ 5.3 & 5.4: ระดับแรงดันและประเภทผู้ใช้ไฟฟ้า"),
    "CH09-07": ("images/diagrams/exam_p25_prob51_52_power_plants.png", "Exam Problem 5.1 & 5.2: Major Power Plants in Thailand", "ข้อสอบจริง ข้อ 5.1 & 5.2: โรงไฟฟ้าสำคัญในไทย"),
    "CH09-09": ("images/diagrams/exam_p24_prob55_56_energy_saving.png", "Exam Problem 5.5: Household Energy Conservation", "ข้อสอบจริง ข้อ 5.5: การอนุรักษ์พลังงานในบ้านเรือน"),
    "CH09-10": ("images/diagrams/exam_p24_prob55_56_energy_saving.png", "Exam Problem 5.6: Industrial Factory Energy Conservation", "ข้อสอบจริง ข้อ 5.6: การอนุรักษ์พลังงานในโรงงานอุตสาหกรรม"),
    "CH10-03": ("images/diagrams/hw10_p2_ammeter_pure.png", "Homework 10 Problem 2: DC Ammeter Design with and without CDRX", "การบ้านบทที่ 10 ข้อ 2: การออกแบบดีซีแอมมิเตอร์ทั้งมีและไม่มี CDRX"),
    "CH10-04": ("images/diagrams/hw10_p3_voltmeter_pure.png", "Homework 10 Problem 3: DC Voltmeter Design with and without CDRX", "การบ้านบทที่ 10 ข้อ 3: การออกแบบดีซีโวลต์มิเตอร์ทั้งมีและไม่มี CDRX"),
    "CH10-05": ("images/diagrams/exam_p29_ammeter_cdrx.png", "Exam Problem 3.5: Multi-Range DC Ammeter with CDRX", "ข้อสอบจริง ข้อ 3.5: ดีซีแอมมิเตอร์หลายย่านวัดพร้อมความต้านทานแดมปิ้งวิกฤต CDRX"),
    "CH10-06": ("images/diagrams/multirange_ammeter.png", "Ayrton / Universal Shunt Multi-Range Ammeter Circuit", "วงจรชันต์ของแอร์ตัน (Ayrton Shunt) สำหรับแอมมิเตอร์"),
    "CH10-07": ("images/diagrams/multirange_voltmeter.png", "Tapped Series Multiplier Multi-Range Voltmeter", "วงจรตัวคูณแรงดันแบบอนุกรมเรียงลำดับสำหรับโวลต์มิเตอร์"),
    "CH10-08": ("images/diagrams/ohmmeter_circuits.png", "Series-Type Ohmmeter Circuit & Scale", "วงจรและสเกลหน้าปัดของโอห์มมิเตอร์แบบอนุกรม"),
    "CH10-09": ("images/diagrams/ohmmeter_circuits.png", "Shunt-Type Ohmmeter Circuit & Scale", "วงจรและสเกลหน้าปัดของโอห์มมิเตอร์แบบขนาน"),
    "CH10-12": ("images/diagrams/exam_p30_wattmeter.png", "Exam Problem 3.4: Electrodynamometer Wattmeter Connection to Load", "ข้อสอบจริง ข้อ 3.4: การต่อขดลวดวัตต์มิเตอร์เข้ากับโหลด"),
    "SEMI-01": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.1: Intrinsic & Extrinsic Semiconductors", "ข้อสอบจริง ข้อ 4.1: สารกึ่งตัวนำ N-Type และ P-Type"),
    "SEMI-02": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.2: MOSFET Definition & Characteristics", "ข้อสอบจริง ข้อ 4.2: โครงสร้างและคุณสมบัติของมอสเฟต (MOSFET)"),
    "SEMI-03": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.3: Integrated Circuit (IC) Significance", "ข้อสอบจริง ข้อ 4.3: วงจรรวม (IC) และประโยชน์ต่อระบบอิเล็กทรอนิกส์"),
    "SEMI-04": ("images/diagrams/diode_circuits.png", "PN Junction Diode Characteristic Curve & Biasing", "กราฟคุณลักษณะและการไบแอสไดโอดรอยต่อ PN"),
    "SEMI-06": ("images/diagrams/diode_circuits.png", "Zener Diode Voltage Regulator Circuit", "วงจรควบคุมแรงดันคงที่ด้วยซีเนอร์ไดโอด (Zener Voltage Regulator)"),
    "SEMI-08": ("images/diagrams/transistor_bjt_mosfet.png", "BJT Transistor Terminal Currents & Alpha-Beta Relations", "กระแสขาใช้งานและอัตราขยายในทรานซิสเตอร์ BJT"),
    "SEMI-10": ("images/diagrams/transistor_bjt_mosfet.png", "Transistor BJT vs MOSFET Comparison", "การเปรียบเทียบเชิงโครงสร้างและการทำงานระหว่าง BJT และ MOSFET")
}

sanitized_cards = []
for c in all_cards:
    card_id = c["id"]
    
    # Strip embedded HTML from questions and solutions
    q_en = strip_embedded_html(c.get("question", ""))
    q_th = strip_embedded_html(c.get("question_th", ""))
    sol_en = strip_embedded_html(c.get("solution", ""))
    sol_th = strip_embedded_html(c.get("solution_th", ""))
    formula = strip_embedded_html(c.get("formula", ""))
    pit_en = strip_embedded_html(c.get("pitfall", ""))
    pit_th = strip_embedded_html(c.get("pitfall_th", ""))
    
    # Check diagram mapping
    if card_id in diagram_map:
        diag_path, cap_en, cap_th = diagram_map[card_id]
    else:
        diag_path = c.get("diagram")
        cap_en = c.get("diagram_caption", "Diagram")
        cap_th = c.get("diagram_caption_th", "รูปประกอบโจทย์")
    
    # Validate diagram file if assigned
    if diag_path and not os.path.exists(diag_path):
        print(f"Warning: diagram path does not exist for {card_id}: {diag_path}")
        diag_path = None

    sanitized_cards.append({
        "id": card_id,
        "chapter": c["chapter"],
        "topic": c["topic"],
        "topic_th": c.get("topic_th", c["topic"]),
        "level": c.get("level", "Level 1 (Concept)"),
        "exam_reference": c.get("exam_reference", "Final Exam"),
        "question": q_en,
        "question_th": q_th,
        "formula": formula,
        "solution": sol_en,
        "solution_th": sol_th,
        "pitfall": pit_en,
        "pitfall_th": pit_th,
        "diagram": diag_path,
        "diagram_caption": cap_en if diag_path else None,
        "diagram_caption_th": cap_th if diag_path else None,
        "mastered": False
    })

print(f"Sanitized {len(sanitized_cards)} cards successfully!")

# Ensure data directory exists
os.makedirs("data", exist_ok=True)

# 1. Export JSON
json_path = os.path.join("data", "flashcards.json")
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(sanitized_cards, f, ensure_ascii=False, indent=2)
print(f"Exported JSON to {json_path}")

# 2. Export JS
js_path = os.path.join("data", "flashcards.js")
with open(js_path, "w", encoding="utf-8") as f:
    f.write("// FunE Final Exam Master Flashcards Data (Bilingual TH / ENG)\n")
    f.write("const FLASHCARDS_DATA = ")
    json.dump(sanitized_cards, f, ensure_ascii=False, indent=2)
    f.write(";\n")
print(f"Exported JS to {js_path}")

# 3. Export Excel Master
excel_path = "FunE_Master_Exam_Flashcards_Ch6_10_Master.xlsx"

rows = []
for c in sanitized_cards:
    rows.append({
        "Card ID": c["id"],
        "Chapter": c["chapter"],
        "Topic (EN)": c["topic"],
        "Topic (TH)": c["topic_th"],
        "Difficulty": c["level"],
        "Exam Reference": c["exam_reference"],
        "Question (EN)": c["question"],
        "Question (TH)": c["question_th"],
        "Formula & Derivations": c["formula"],
        "Solution (EN)": c["solution"],
        "Solution (TH)": c["solution_th"],
        "Exam Pitfall (EN)": c["pitfall"],
        "Exam Pitfall (TH)": c["pitfall_th"],
        "Diagram Path": c.get("diagram") or "",
        "Diagram Caption (TH)": c.get("diagram_caption_th") or "",
        "Mastered": "Yes" if c.get("mastered") else "No"
    })

df = pd.DataFrame(rows)
df.to_excel(excel_path, index=False, engine="openpyxl")

# Format Excel with professional styling
wb = openpyxl.load_workbook(excel_path)
ws = wb.active
ws.title = "FunE_Bilingual_Flashcards"

header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for col_num in range(1, len(df.columns) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

row_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

for r_idx in range(2, len(df) + 2):
    fill = row_fill_even if r_idx % 2 == 0 else row_fill_odd
    for c_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.fill = fill
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)
        if c_idx in [1, 2, 5, 14, 16]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

col_widths = {
    "A": 12, "B": 15, "C": 28, "D": 28, "E": 18, "F": 25,
    "G": 42, "H": 42, "I": 35, "J": 45, "K": 45, "L": 32, "M": 32, "N": 25, "O": 30, "P": 10
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

ws.freeze_panes = "A2"
wb.save(excel_path)
print(f"Exported styled bilingual Excel to {excel_path} successfully!")
