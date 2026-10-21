# -*- coding: utf-8 -*-
"""
build_full_78_cards.py
Assembles the complete 78 master flashcards:
- Chapter 6: Three-Phase Circuits (CH06-01 to CH06-10) [10 cards]
- Chapter 7: Electric Machinery & Transformers (CH07-01 to CH07-27) [27 cards]
- Chapter 8: Power Distribution System (CH08-01 to CH08-15) [15 cards]
- Chapter 9: Electrical Measuring Instruments (CH09-01 to CH09-14) [14 cards]
- Chapter 10: Semiconductor Devices & Circuits (CH10-01 to CH10-12) [12 cards]
"""

import json
import re
import os

from ch08_cards_data import ch08_cards
import create_bilingual_part1
import create_bilingual_part2
import create_bilingual_part3

def strip_embedded_html(text):
    if not isinstance(text, str):
        return text
    text = re.sub(r'<div class=[\'\"].*?</div></div>', '', text, flags=re.DOTALL)
    text = re.sub(r'<div class=[\'\"].*?</div>', '', text, flags=re.DOTALL)
    text = text.replace('\x0b', '\\vec')
    return text.strip()

# 1. Chapter 6 (10 cards) & Chapter 7 Machinery (15 cards)
p1_cards = create_bilingual_part1.cards
print(f"Loaded {len(p1_cards)} cards from part 1")

# Extract Ch6 (10 cards) and Ch7 Machinery (15 cards)
ch6_cards = [c for c in p1_cards if c['chapter'] == 'Chapter 6']
ch7_machinery = [c for c in p1_cards if c['chapter'] == 'Chapter 7']

# 2. Extract Transformers from part 2 (formerly CH08-01 to CH08-12)
p2_cards = create_bilingual_part2.full_cards
old_transformer_cards = [c for c in p2_cards if 'Transformer' in c.get('topic', '') or c['id'].startswith('CH08')]

ch7_transformers = []
for idx, old_c in enumerate(old_transformer_cards):
    new_card = dict(old_c)
    new_num = 16 + idx
    new_card['id'] = f"CH07-{new_num:02d}"
    new_card['chapter'] = "Chapter 7"
    # Ensure topic indicates Transformers
    if "Transformer" not in new_card['topic']:
        new_card['topic'] = f"Transformer: {new_card['topic']}"
    ch7_transformers.append(new_card)

print(f"Re-indexed {len(ch7_transformers)} transformer cards to CH07-16..CH07-27")

ch7_all = ch7_machinery + ch7_transformers
print(f"Total Chapter 7 cards: {len(ch7_all)}")

# 3. Chapter 8: Power Distribution System (15 cards from ch08_cards_data)
print(f"Total Chapter 8 cards: {len(ch08_cards)}")

# 4. Chapter 9: Measuring Instruments (formerly CH10-01 to CH10-14 in part 3)
p3_cards = create_bilingual_part3.cards_p3
old_meter_cards = [c for c in p3_cards if c.get('chapter') == 'Chapter 10' or c['id'].startswith('CH10')]

ch9_instruments = []
for idx, old_c in enumerate(old_meter_cards):
    new_card = dict(old_c)
    new_num = idx + 1
    new_card['id'] = f"CH09-{new_num:02d}"
    new_card['chapter'] = "Chapter 9"
    ch9_instruments.append(new_card)

print(f"Re-indexed {len(ch9_instruments)} instrument cards to CH09-01..CH09-14")

# 5. Chapter 10: Semiconductor Devices (formerly SEMI-01 to SEMI-12 in part 3)
old_semi_cards = [c for c in p3_cards if c.get('chapter') == 'Semiconductors' or c['id'].startswith('SEMI')]

ch10_semiconductors = []
for idx, old_c in enumerate(old_semi_cards):
    new_card = dict(old_c)
    new_num = idx + 1
    new_card['id'] = f"CH10-{new_num:02d}"
    new_card['chapter'] = "Chapter 10"
    ch10_semiconductors.append(new_card)

print(f"Re-indexed {len(ch10_semiconductors)} semiconductor cards to CH10-01..CH10-12")

# Combine all 78 master cards
final_cards_raw = ch6_cards + ch7_all + ch08_cards + ch9_instruments + ch10_semiconductors
print(f"\nTOTAL MASTER CARDS: {len(final_cards_raw)}")

# Diagram mapping dictionary
diagram_map = {
    # Chapter 6
    "CH06-04": ("images/diagrams/hw6_fig1_yy_balanced.png", "Homework 6 Figure 1: Balanced Y-Y System", "การบ้านบทที่ 6 รูปที่ 1: ระบบ 3 เฟสแบบ Y-Y สมดุล"),
    "CH06-05": ("images/diagrams/hw6_fig2_yy_unbalanced.png", "Homework 6 Figure 2: Unbalanced Y-Y 4-Wire with Neutral Line", "การบ้านบทที่ 6 รูปที่ 2: ระบบ 3 เฟสแบบ Y-Y ไม่สมดุล มีสายนิวทรัล"),
    "CH06-06": ("images/diagrams/hw6_fig3_yy_neutral_z.png", "Homework 6 Figure 3: Unbalanced Y-Y 3-Wire System", "การบ้านบทที่ 6 รูปที่ 3: ระบบ 3 เฟสแบบ Y-Y ไม่สมดุล ไม่มีสายนิวทรัล"),
    
    # Chapter 7 Machinery
    "CH07-01": ("images/diagrams/exam_p28_prob41_right_hand.png", "Exam Problem 4.1: Right-Hand Grip Rule", "ข้อสอบจริง ข้อ 4.1: กฎมือขวาของสนามแม่เหล็กรอบเส้นลวด"),
    "CH07-02": ("images/diagrams/exam_p28_prob42_motor_force.png", "Exam Problem 4.2: DC Motor Magnetic Field & Force (Lorentz)", "ข้อสอบจริง ข้อ 4.2: สนามแม่เหล็กและแรงลอเรนซ์ในมอเตอร์ไฟฟ้ากระแสตรง"),
    "CH07-03": ("images/diagrams/exam_p28_prob43_motor_poles.png", "Exam Problem 4.3: Motor Magnetic Poles (N & S)", "ข้อสอบจริง ข้อ 4.3: ขั้วแม่เหล็ก N และ S ของมอเตอร์ไฟฟ้า"),
    "CH07-04": ("images/diagrams/hw7_p4_shunt_circuit.png", "Homework 7 Problem 4: DC Shunt Motor Equivalent Circuit", "การบ้านบทที่ 7 ข้อ 4: วงจรสมมูลมอเตอร์ไฟฟ้ากระแสตรงแบบชันต์ (DC Shunt Motor)"),
    "CH07-05": ("images/diagrams/hw7_p5_series_circuit.png", "Homework 7 Problem 5: DC Series Motor Equivalent Circuit", "การบ้านบทที่ 7 ข้อ 5: วงจรสมมูลมอเตอร์ไฟฟ้ากระแสตรงแบบอนุกรม (DC Series Motor)"),
    "CH07-06": ("images/diagrams/exam_p27_prob44_induction_motor.png", "Exam Problem 4.4: 3-Phase Induction Motor (Squirrel Cage)", "ข้อสอบจริง ข้อ 4.4: โครงสร้างมอเตอร์เหนี่ยวนำ 3 เฟส (โรเตอร์กรงกระรอก)"),
    "CH07-07": ("images/diagrams/exam_p26_prob45_generator.png", "Exam Problem 4.5: DC Generator Principle & Split-Ring Commutator", "ข้อสอบจริง ข้อ 4.5: เครื่องกำเนิดไฟฟ้ากระแสตรงและคอมมิวเทเตอร์"),
    
    # Chapter 7 Transformers
    "CH07-16": ("images/diagrams/exam_p32_transformer_core.png", "Exam Problem 1.2: Transformer Core & Dual Windings", "ข้อสอบจริง ข้อ 1.2: แกนเหล็กและขดลวดปฐมภูมิ-ทุติยภูมิของหม้อแปลง"),
    "CH07-17": ("images/diagrams/exam_p31_prob16_transformer_calc.png", "Exam Problem 1.6: 25 kVA Transformer Full Calculation", "ข้อสอบจริง ข้อ 1.6: การคำนวณหม้อแปลง 25 kVA ครบวงจร"),
    "CH07-23": ("images/diagrams/transformer_audio_matching.png", "Audio Transformer Maximum Power Transfer Impedance Matching", "หม้อแปลงขยายเสียงสำหรับปรับอิมพีแดนซ์ให้เกิดการส่งถ่ายกำลังสูงสุด"),
    
    # Chapter 8 Power Distribution System
    "CH08-01": ("images/ch08_slide_7.png", "Slide 7: EGAT 4 High-Voltage Transmission Levels & Power Systems Structure", "สไลด์ที่ 7: โครงสร้าง 4 ขั้นตอนของระบบไฟฟ้า และระดับแรงดันสายส่ง 4 ระดับของ กฟผ."),
    "CH08-06": ("images/ch08_slide_20.png", "Slide 20: Three Hydraulic Turbine Types (Kaplan, Francis, Pelton)", "สไลด์ที่ 20: กังหันน้ำ 3 ชนิด (แคปแลน, ฟรานซิส, เพลตัน)"),
    "CH08-07": ("images/ch08_slide_35.png", "Slide 35: Components of a Solar Power Plant", "สไลด์ที่ 35: องค์ประกอบ 5 ส่วนของโรงไฟฟ้าพลังงานแสงอาทิตย์"),
    "CH08-08": ("images/ch08_slide_42.png", "Slide 42: Single Line Diagram of MEA Distribution System (69 kV / 24 kV)", "สไลด์ที่ 42: ไดอะแกรมเส้นเดี่ยว (Single Line Diagram) ระบบจำหน่าย กฟน. 69 kV / 24 kV"),
    "CH08-09": ("images/ch08_slide_45.png", "Slide 45 & 46: PEA 22 kV Solidly Grounded vs 33 kV Multigrounded System", "สไลด์ที่ 45 และ 46: ระบบจำหน่าย 22 kV ต่อลงดินโดยตรง เทียบกับ 33 kV มัลติกราวด์ กฟภ."),
    "CH08-10": ("images/ch08_slide_47.png", "Slide 47 & 48: Radial Line vs Tie-Line Distribution Topologies", "สไลด์ที่ 47 และ 48: สายป้อนแบบรัศมี เทียบกับ แบบเชื่อมโยงระหว่างสถานี กฟภ."),
    "CH08-11": ("images/ch08_slide_50.png", "Slide 50: PEA 3-Phase 4-Wire Low Voltage Distribution System (400/230 V)", "สไลด์ที่ 50: ระบบจำหน่ายแรงดันต่ำ 3 เฟส 4 สาย 400/230 V ของ กฟภ."),
    "CH08-12": ("images/ch08_slide_54.png", "Slide 54: Medium Voltage Distribution Connection & Boundary of Responsibility", "สไลด์ที่ 54: การต่อระบบจำหน่ายแรงดันปานกลางและขอบเขตความรับผิดชอบ"),
    "CH08-13": ("images/ch08_slide_58.png", "Slide 58 & 59: Elements of Substations and KKU Indoor Substation", "สไลด์ที่ 58 และ 59: อุปกรณ์สถานีไฟฟ้าย่อย และสถานีย่อยในร่ม มหาวิทยาลัยขอนแก่น (KKU GIS)"),
    "CH08-14": ("images/diagrams/transmission_towers.png", "High-Voltage Transmission Towers (Suspension, Tension, Transposition)", "โครงสร้างเสาสายส่งไฟฟ้าแรงสูง (เสาแขวน, เสารั้ง, เสาสลับเฟส)"),
    "CH08-15": ("images/diagrams/exam_p24_prob55_56_energy_saving.png", "Final Exam Problems 5.5 & 5.6: Energy Conservation in Home and Factory", "ข้อสอบจริง ข้อ 5.5 และ 5.6: การอนุรักษ์พลังงานในบ้านและโรงงานอุตสาหกรรม"),

    # Chapter 9 Measuring Instruments
    "CH09-03": ("images/diagrams/hw10_p2_ammeter_pure.png", "Homework 10 Problem 2: DC Ammeter Design with and without CDRX", "การบ้านบทที่ 10 ข้อ 2: การออกแบบดีซีแอมมิเตอร์ทั้งมีและไม่มี CDRX"),
    "CH09-04": ("images/diagrams/hw10_p3_voltmeter_pure.png", "Homework 10 Problem 3: DC Voltmeter Design with and without CDRX", "การบ้านบทที่ 10 ข้อ 3: การออกแบบดีซีโวลต์มิเตอร์ทั้งมีและไม่มี CDRX"),
    "CH09-05": ("images/diagrams/exam_p29_ammeter_cdrx.png", "Exam Problem 3.5: Multi-Range DC Ammeter with CDRX", "ข้อสอบจริง ข้อ 3.5: ดีซีแอมมิเตอร์หลายย่านวัดพร้อมความต้านทานแดมปิ้งวิกฤต CDRX"),
    "CH09-06": ("images/diagrams/multirange_ammeter.png", "Ayrton / Universal Shunt Multi-Range Ammeter Circuit", "วงจรชันต์ของแอร์ตัน (Ayrton Shunt) สำหรับแอมมิเตอร์"),
    "CH09-07": ("images/diagrams/multirange_voltmeter.png", "Tapped Series Multiplier Multi-Range Voltmeter", "วงจรตัวคูณแรงดันแบบอนุกรมเรียงลำดับสำหรับโวลต์มิเตอร์"),
    "CH09-08": ("images/diagrams/ohmmeter_circuits.png", "Series-Type Ohmmeter Circuit & Scale", "วงจรและสเกลหน้าปัดของโอห์มมิเตอร์แบบอนุกรม"),
    "CH09-09": ("images/diagrams/ohmmeter_circuits.png", "Shunt-Type Ohmmeter Circuit & Scale", "วงจรและสเกลหน้าปัดของโอห์มมิเตอร์แบบขนาน"),
    "CH09-12": ("images/diagrams/exam_p30_wattmeter.png", "Exam Problem 3.4: Electrodynamometer Wattmeter Connection to Load", "ข้อสอบจริง ข้อ 3.4: การต่อขดลวดวัตต์มิเตอร์เข้ากับโหลด"),

    # Chapter 10 Semiconductors
    "CH10-01": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.1: Intrinsic & Extrinsic Semiconductors", "ข้อสอบจริง ข้อ 4.1: สารกึ่งตัวนำ N-Type และ P-Type"),
    "CH10-02": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.2: MOSFET Definition & Characteristics", "ข้อสอบจริง ข้อ 4.2: โครงสร้างและคุณสมบัติของมอสเฟต (MOSFET)"),
    "CH10-03": ("images/diagrams/exam_p24_prob4_semiconductors.png", "Exam Problem 4.3: Integrated Circuit (IC) Significance", "ข้อสอบจริง ข้อ 4.3: วงจรรวม (IC) และประโยชน์ต่อระบบอิเล็กทรอนิกส์"),
    "CH10-04": ("images/diagrams/diode_circuits.png", "PN Junction Diode Characteristic Curve & Biasing", "กราฟคุณลักษณะและการไบแอสไดโอดรอยต่อ PN"),
    "CH10-06": ("images/diagrams/diode_circuits.png", "Zener Diode Voltage Regulator Circuit", "วงจรควบคุมแรงดันคงที่ด้วยซีเนอร์ไดโอด (Zener Voltage Regulator)"),
    "CH10-08": ("images/diagrams/transistor_bjt_mosfet.png", "BJT Transistor Terminal Currents & Alpha-Beta Relations", "กระแสขาใช้งานและอัตราขยายในทรานซิสเตอร์ BJT"),
    "CH10-10": ("images/diagrams/transistor_bjt_mosfet.png", "Transistor BJT vs MOSFET Comparison", "การเปรียบเทียบเชิงโครงสร้างและการทำงานระหว่าง BJT และ MOSFET")
}

final_clean_cards = []
for c in final_cards_raw:
    cid = c['id']
    q_en = strip_embedded_html(c.get('question', ''))
    q_th = strip_embedded_html(c.get('question_th', ''))
    sol_en = strip_embedded_html(c.get('solution', ''))
    sol_th = strip_embedded_html(c.get('solution_th', ''))
    formula = strip_embedded_html(c.get('formula', ''))
    pit_en = strip_embedded_html(c.get('pitfall', ''))
    pit_th = strip_embedded_html(c.get('pitfall_th', ''))
    
    if cid in diagram_map:
        diag_path, cap_en, cap_th = diagram_map[cid]
    else:
        diag_path = c.get('diagram')
        cap_en = c.get('diagram_caption', 'Diagram')
        cap_th = c.get('diagram_caption_th', 'รูปภาพประกอบ')
    
    # Check diagram path
    if diag_path and not os.path.exists(os.path.join(r"c:\Users\thana\Desktop\Private-Flashcard", diag_path)):
        print(f"Notice: diag_path missing: {diag_path} for {cid}")
    
    card_obj = {
        "id": cid,
        "chapter": c["chapter"],
        "topic": c["topic"],
        "topic_th": c["topic_th"],
        "level": c.get("level", "Level 1 (Concept)"),
        "exam_reference": c.get("exam_reference", ""),
        "question": q_en,
        "question_th": q_th,
        "formula": formula,
        "solution": sol_en,
        "solution_th": sol_th,
        "pitfall": pit_en,
        "pitfall_th": pit_th,
        "diagram": diag_path,
        "diagram_caption": cap_en,
        "diagram_caption_th": cap_th,
        "mastered": False
    }
    final_clean_cards.append(card_obj)

print(f"\nFinal cleaned cards count: {len(final_clean_cards)}")

# Save to data/flashcards.json
out_json_path = r"c:\Users\thana\Desktop\Private-Flashcard\data\flashcards.json"
with open(out_json_path, 'w', encoding='utf-8') as f:
    json.dump(final_clean_cards, f, ensure_ascii=False, indent=2)
print(f"Saved {out_json_path}")

# Save to data/flashcards.js
out_js_path = r"c:\Users\thana\Desktop\Private-Flashcard\data\flashcards.js"
with open(out_js_path, 'w', encoding='utf-8') as f:
    f.write("const flashcardsData = " + json.dumps(final_clean_cards, ensure_ascii=False, indent=2) + ";\n")
print(f"Saved {out_js_path}")

# Also copy to Fund E N'Manow/data
fund_e_data = r"c:\Users\thana\Desktop\Fund E N'Manow\data"
os.makedirs(fund_e_data, exist_ok=True)
with open(os.path.join(fund_e_data, 'flashcards.json'), 'w', encoding='utf-8') as f:
    json.dump(final_clean_cards, f, ensure_ascii=False, indent=2)
with open(os.path.join(fund_e_data, 'flashcards.js'), 'w', encoding='utf-8') as f:
    f.write("const flashcardsData = " + json.dumps(final_clean_cards, ensure_ascii=False, indent=2) + ";\n")
print("Synced JSON/JS to Fund E N'Manow/data")

# 3. Export Excel Master
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import shutil

excel_path = "FunE_Master_Exam_Flashcards_Ch6_10_Master.xlsx"

rows = []
for c in final_clean_cards:
    rows.append({
        "Card ID": c["id"],
        "Chapter": c["chapter"],
        "Topic (EN)": c["topic"],
        "Topic (TH)": c["topic_th"],
        "Difficulty": c["level"],
        "Exam Reference": c["exam_reference"],
        "Question (EN)": c["question"],
        "Question (TH)": c["question_th"],
        "Formula & Derivations": c["formula"],
        "Solution (EN)": c["solution"],
        "Solution (TH)": c["solution_th"],
        "Exam Pitfall (EN)": c["pitfall"],
        "Exam Pitfall (TH)": c["pitfall_th"],
        "Diagram Path": c.get("diagram") or "",
        "Diagram Caption (TH)": c.get("diagram_caption_th") or "",
        "Mastered": "Yes" if c.get("mastered") else "No"
    })

df = pd.DataFrame(rows)
df.to_excel(excel_path, index=False, engine="openpyxl")

wb = openpyxl.load_workbook(excel_path)
ws = wb.active
ws.title = "FunE_Bilingual_Flashcards"

header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for col_num in range(1, len(df.columns) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

row_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

for r_idx in range(2, len(df) + 2):
    fill = row_fill_even if r_idx % 2 == 0 else row_fill_odd
    for c_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.fill = fill
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)
        if c_idx in [1, 2, 5, 14, 16]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

col_widths = {
    "A": 12, "B": 15, "C": 28, "D": 28, "E": 18, "F": 25,
    "G": 42, "H": 42, "I": 35, "J": 45, "K": 45, "L": 32, "M": 32, "N": 25, "O": 30, "P": 10
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

ws.freeze_panes = "A2"
wb.save(excel_path)
print(f"Exported styled bilingual Excel to {excel_path} successfully!")

# Copy Excel to Fund E N'Manow
fund_e_excel = r"c:\Users\thana\Desktop\Fund E N'Manow\FunE_Master_Exam_Flashcards_Ch6_10_Master.xlsx"
try:
    shutil.copy2(excel_path, fund_e_excel)
    print(f"Synced Excel to {fund_e_excel}")
except PermissionError:
    print(f"Notice: {fund_e_excel} is currently open in Excel. Saved locally to {excel_path}.")

try:
    shutil.copy2(excel_path, "FunE_Master_Exam_Flashcards_Ch6_10.xlsx")
    shutil.copy2(excel_path, r"c:\Users\thana\Desktop\Fund E N'Manow\FunE_Master_Exam_Flashcards_Ch6_10.xlsx")
except Exception as e:
    print(f"Notice: Secondary copy skipped: {e}")

