import os
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Helper for embedding diagram images in HTML
def make_img_html(img_name, caption=""):
    cap_html = f"<div class='text-xs text-center text-slate-400 mt-1 mb-2'>{caption}</div>" if caption else ""
    return f"<div class='my-3 text-center'><img src='images/diagrams/{img_name}' alt='{caption}' class='rounded-lg border border-slate-300 dark:border-slate-700 max-h-60 mx-auto shadow-sm object-contain' />{cap_html}</div>"

flashcards = [
    # =========================================================================
    # CHAPTER 6: THREE-PHASE CIRCUITS (10 CARDS)
    # =========================================================================
    {
        "id": "CH06-01",
        "chapter": "Chapter 6",
        "topic": "Star (Y) vs Delta (Δ) Voltage & Current Relations",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides & Formula Sheet",
        "question": "In a balanced 3-phase positive-sequence ($abc$) system, state the exact relationships between line quantities ($V_L, I_L$) and phase quantities ($V_p, I_p$) for both Y-connection and $\\Delta$-connection.",
        "formula": "Y: $V_L = \\sqrt{3}V_p\\angle +30^\\circ, I_L = I_p$\n$\\Delta$: $V_L = V_p, I_L = \\sqrt{3}I_p\\angle -30^\\circ$",
        "solution": "1. Star (Y) Connection:\n- Line voltage leads phase voltage by $30^\\circ$: $V_{ab} = \\sqrt{3}V_{an}\\angle +30^\\circ$\n- Line current equals phase current: $I_L = I_p$\n\n2. Delta (Δ) Connection:\n- Line voltage equals phase voltage: $V_L = V_p$\n- Line current lags phase current by $30^\\circ$: $I_a = \\sqrt{3}I_{ab}\\angle -30^\\circ$",
        "pitfall": "In negative sequence (acb), line voltage leads by -30° (or lags by 30°). Remember: in Y, currents are identical; in Δ, voltages are identical.",
        "mastered": False
    },
    {
        "id": "CH06-02",
        "chapter": "Chapter 6",
        "topic": "Delta to Wye (Δ ↔ Y) Impedance Conversion",
        "level": "Level 1 (Concept)",
        "exam_reference": "HW Ch.6 Problem 1 & 5",
        "question": "If a balanced $3\\phi$ load is connected in $\\Delta$ with impedance $Z_\\Delta = 60 + j30\\ \\Omega$, what is the equivalent impedance per phase when converted to a Wye (Y) connection ($Z_Y$)?",
        "formula": "$Z_Y = \\frac{Z_\\Delta}{3}$",
        "solution": "$$Z_Y = \\frac{Z_\\Delta}{3} = \\frac{60 + j30}{3} = 20 + j10\\ \\Omega$$\nConversely, if converting from Wye to Delta: $Z_\\Delta = 3 Z_Y$.",
        "pitfall": "Multiplying by 3 instead of dividing is a classic mistake! A Delta load draws 3 times more current than a Wye load of the same impedance value, so equivalent Wye impedance must be 3 times smaller.",
        "mastered": False
    },
    {
        "id": "CH06-03",
        "chapter": "Chapter 6",
        "topic": "Three-Phase Power Formulas & Complex Power",
        "level": "Level 1 (Concept)",
        "exam_reference": "Formula Sheet & Lecture Slides",
        "question": "State the formulas for balanced $3\\phi$ Real Power ($P_{3\\phi}$), Reactive Power ($Q_{3\\phi}$), Apparent Power ($|S_{3\\phi}|$), and Complex Power ($S_{3\\phi}$) in terms of line values ($V_L, I_L$).",
        "formula": "$P_{3\\phi} = \\sqrt{3}V_L I_L \\cos\\theta\\quad [\\text{W}]$\n$Q_{3\\phi} = \\sqrt{3}V_L I_L \\sin\\theta\\quad [\\text{VAR}]$\n$|S_{3\\phi}| = \\sqrt{3}V_L I_L\\quad [\\text{VA}]$\n$S_{3\\phi} = \\sqrt{3}V_L I_L^* = P + jQ$",
        "solution": "When using Line values ($V_L, I_L$), always include the factor $\\sqrt{3}$.\nWhen using Phase values ($V_p, I_p$), use factor $3$:\n- $P = 3 V_p I_p \\cos\\theta$\n- $Q = 3 V_p I_p \\sin\\theta$\n- $S = 3 V_p I_p^*$\nwhere $\\theta = \\angle Z_p$ is the load phase impedance angle, and $PF = \\cos\\theta$.",
        "pitfall": "Never mix $\\sqrt{3}$ with phase quantities! It is either $\\sqrt{3}V_L I_L$ or $3 V_p I_p$.",
        "mastered": False
    },
    {
        "id": "CH06-04",
        "chapter": "Chapter 6",
        "topic": "Balanced 4-Wire Y-Y System Full Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.6 Problem 1 (Page 1)",
        "question": "In the balanced Y-Y circuit shown below, the source phase voltages are $V_{an} = 220\\angle 0^\\circ\\text{ V}$, $V_{bn} = 220\\angle -120^\\circ\\text{ V}$, $V_{cn} = 220\\angle -240^\\circ\\text{ V}$. Line impedance is $Z_{\\text{line}} = 10 + j5\\ \\Omega$. Balanced load impedance is $Z_L = 20 + j50\\ \\Omega$. Find: (1) Line voltages $V_{ab}, V_{bc}, V_{ca}$, (2) Line currents $I_{aA}, I_{bB}, I_{cC}$, (3) Neutral current $I_{Nn}$, (4) Total $3\\phi$ complex power $S$.\n" + make_img_html('hw6_fig1_yy_balanced.png', 'Fig. 1: Balanced Y-Y System with Line Impedance'),
        "formula": "$V_{ab} = \\sqrt{3}V_{an}\\angle +30^\\circ$\n$I_{aA} = \\frac{V_{an}}{Z_{\\text{line}} + Z_L}$\n$I_{Nn} = 0$\n$S_{3\\phi} = 3 V_{an} I_{aA}^*$",
        "solution": "1. Line Voltages:\n- $V_{ab} = \\sqrt{3}(220)\\angle (0^\\circ + 30^\\circ) = 381.05\\angle 30^\\circ\\text{ V}$\n- $V_{bc} = 381.05\\angle -90^\\circ\\text{ V}$\n- $V_{ca} = 381.05\\angle -210^\\circ\\text{ V}$\n\n2. Per-phase total impedance:\n$Z_{\\text{total}} = (10 + j5) + (20 + j50) = 30 + j55\\ \\Omega = 62.65\\angle 61.39^\\circ\\ \\Omega$\n\n3. Line Currents:\n- $I_{aA} = \\frac{220\\angle 0^\\circ}{62.65\\angle 61.39^\\circ} = 3.51\\angle -61.39^\\circ\\text{ A}$\n- $I_{bB} = 3.51\\angle -181.39^\\circ\\text{ A}$\n- $I_{cC} = 3.51\\angle 58.61^\\circ\\text{ A}$\n\n4. Neutral Current:\nIn a balanced system, $I_{Nn} = -(I_{aA} + I_{bB} + I_{cC}) = 0\\text{ A}$. $V_{Nn} = 0\\text{ V}$.\n\n5. Total $3\\phi$ Power:\n$S = 3 \\times 220 \\times 3.51\\angle +61.39^\\circ = 2,316.6\\angle 61.39^\\circ\\text{ VA} = 1,109\\text{ W} + j2,034\\text{ VAR}$.",
        "pitfall": "In a balanced system, neutral voltage shift $V_{Nn}$ is strictly ZERO, and neutral current $I_{Nn} = 0$.",
        "mastered": False
    },
    {
        "id": "CH06-05",
        "chapter": "Chapter 6",
        "topic": "Unbalanced 3-Wire Y-Y System (Neutral Shift)",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.6 Problem 2 (Page 2)",
        "question": "In the 3-wire Y-Y system below (NO neutral wire), source is $220\\text{ V}$ balanced, $Z_{\\text{line}} = 10 + j5\\ \\Omega$, but load is UNBALANCED: $Z_A = 20 + j50\\ \\Omega, Z_B = 10 + j50\\ \\Omega, Z_C = 20 + j20\\ \\Omega$. How is the neutral voltage shift $V_{Nn}$ calculated, and why do load phase voltages become unbalanced?\n" + make_img_html('hw6_fig2_yy_unbalanced.png', 'Fig. 2: Unbalanced 3-Wire Y-Y System without Neutral'),
        "formula": "$V_{Nn} = \\frac{V_{an} Y_A + V_{bn} Y_B + V_{cn} Y_C}{Y_A + Y_B + Y_C}$ (Millman's Theorem)",
        "solution": "1. Millman's Theorem for Neutral Shift:\nDefine total branch admittance $Y_A = 1/(Z_{\\text{line}} + Z_A)$, $Y_B = 1/(Z_{\\text{line}} + Z_B)$, $Y_C = 1/(Z_{\\text{line}} + Z_C)$.\n$$V_{Nn} = \\frac{V_{an}Y_A + V_{bn}Y_B + V_{cn}Y_C}{Y_A + Y_B + Y_C} \\neq 0$$\n\n2. Line Currents with Neutral Shift:\n- $I_{aA} = (V_{an} - V_{Nn}) Y_A$\n- $I_{bB} = (V_{bn} - V_{Nn}) Y_B$\n- $I_{cC} = (V_{cn} - V_{Nn}) Y_C$\nSince no neutral wire exists, $I_{aA} + I_{bB} + I_{cC} = 0$.\n\n3. Impact on Load Phase Voltages:\nLoad phase voltage $V_{AN} = V_{an} - I_{aA} Z_{\\text{line}} - V_{Nn}$. Because $V_{Nn} \\neq 0$, phase voltages across the load become severely distorted and unbalanced.",
        "pitfall": "Without a neutral wire, the neutral point floats ($V_{Nn} \\neq 0$), causing dangerous overvoltage on lighter-loaded phases!",
        "mastered": False
    },
    {
        "id": "CH06-06",
        "chapter": "Chapter 6",
        "topic": "Unbalanced Y-Y System with Neutral Line Impedance",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.6 Problem 3 (Page 2)",
        "question": "In the 4-wire Y-Y system below, an unbalanced load ($Z_A \\neq Z_B \\neq Z_C$) is connected with a neutral return wire having finite impedance $Z_{nN}$. Write the exact node-voltage equation to solve for $V_{Nn}$ and neutral current $I_{Nn}$.\n" + make_img_html('hw6_fig3_yy_neutral_z.png', 'Fig. 3: Unbalanced 4-Wire Y-Y System with Neutral Impedance'),
        "formula": "$V_{Nn} = \\frac{V_{an} Y_A + V_{bn} Y_B + V_{cn} Y_C}{Y_A + Y_B + Y_C + Y_n}$\n$I_{Nn} = V_{Nn} Y_n = \\frac{V_{Nn}}{Z_{nN}}$",
        "solution": "1. Node-Voltage Equation at node N:\nBy KCL at load neutral node N:\n$$\\frac{V_{Nn} - V_{an}}{Z_A + Z_{\\text{line}}} + \\frac{V_{Nn} - V_{bn}}{Z_B + Z_{\\text{line}}} + \\frac{V_{Nn} - V_{cn}}{Z_C + Z_{\\text{line}}} + \\frac{V_{Nn}}{Z_{nN}} = 0$$\n\n2. Solving for $V_{Nn}$:\n$$V_{Nn} = \\frac{V_{an}Y_A' + V_{bn}Y_B' + V_{cn}Y_C'}{Y_A' + Y_B' + Y_C' + Y_n}$$\nwhere $Y_k' = 1/(Z_k + Z_{\\text{line}})$ and $Y_n = 1/Z_{nN}$.\n\n3. Neutral Return Current:\n$$I_{Nn} = \\frac{V_{Nn}}{Z_{nN}} = -(I_{aA} + I_{bB} + I_{cC})$$.",
        "pitfall": "Do not omit $Y_n = 1/Z_{nN}$ in the denominator of Millman's formula! As $Z_{nN} \\to 0$, $V_{Nn} \\to 0$ (solid ground).",
        "mastered": False
    },
    {
        "id": "CH06-07",
        "chapter": "Chapter 6",
        "topic": "Y-Connected Induction Motor Load Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.6 Problem 4 (Page 3)",
        "question": "A $3\\phi$, $12\\text{-hp}$ induction motor is Y-connected to a $3\\phi$ supply with line-to-line voltage $V_L = 380\\text{ V}, 50\\text{ Hz}$. The motor has $\\eta = 90\\%$ efficiency and a lagging power factor of $0.85$. Calculate: (1) Input real power ($P_{\\text{in}}$), (2) Line current ($I_L$), (3) Per-phase impedance magnitude ($|Z|$), resistance ($R$), and inductance ($L$). ($1\\text{ hp} = 746\\text{ W}$)",
        "formula": "$P_{\\text{out}} = 12 \\times 746\\text{ W}$\n$P_{\\text{in}} = P_{\\text{out}} / \\eta$\n$I_L = \\frac{P_{\\text{in}}}{\\sqrt{3} V_L \\cos\\theta}$\n$V_p = V_L / \\sqrt{3}$\n$Z = V_p / I_p$",
        "solution": "1. Input Power:\n$P_{\\text{out}} = 12 \\times 746 = 8,952\\text{ W}$\n$P_{\\text{in}} = 8,952 / 0.90 = 9,946.67\\text{ W}$ ($9.95\\text{ kW}$)\n\n2. Line Current:\n$$I_L = \\frac{9,946.67}{\\sqrt{3} \\times 380 \\times 0.85} = 17.78\\text{ A}$$\n\n3. Phase Impedance (Y-connection: $V_p = 380/\\sqrt{3} = 219.39\\text{ V}, I_p = I_L = 17.78\\text{ A}$):\n- $|Z| = 219.39 / 17.78 = 12.34\\ \\Omega$\n- $\\theta = \\arccos(0.85) = 31.79^\\circ$\n- $Z = 12.34\\angle 31.79^\\circ = 10.49 + j6.50\\ \\Omega$\n- Resistance: $R = 10.49\\ \\Omega$\n- Inductance: $X_L = 2\\pi f L = 6.50 \\implies L = 6.50 / (2\\pi \\times 50) = 20.7\\text{ mH}$.",
        "pitfall": "Always convert output mechanical horsepower to electrical input power using efficiency ($P_{\\text{in}} = P_{\\text{out}}/\\eta$) BEFORE calculating current!",
        "mastered": False
    },
    {
        "id": "CH06-08",
        "chapter": "Chapter 6",
        "topic": "Delta-Connected Motor Load Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.6 Problem 5 & Lecture Slides",
        "question": "A $3\\phi$, $10\\text{-hp}$ motor is $\\Delta$-connected to a $3\\phi$ supply with $V_L = 380\\text{ V}, 50\\text{ Hz}$, $\\eta = 85\\%$, and $PF = 0.8$ lagging. Find: (1) Line current ($I_L$), (2) Phase current ($I_p$), (3) Per-phase $\\Delta$ impedance ($Z_\\Delta$).",
        "formula": "$P_{\\text{in}} = (10 \\times 746) / 0.85 = 8,776.47\\text{ W}$\n$I_L = \\frac{P_{\\text{in}}}{\\sqrt{3} V_L \\cos\\theta}$\n$I_p = I_L / \\sqrt{3}, V_p = V_L = 380\\text{ V}$\n$Z_\\Delta = V_p / I_p$",
        "solution": "1. Input Power: $P_{\\text{in}} = 7,460 / 0.85 = 8,776.47\\text{ W}$\n\n2. Line Current:\n$$I_L = \\frac{8,776.47}{\\sqrt{3} \\times 380 \\times 0.8} = 16.67\\text{ A}$$\n\n3. Phase Quantities & Impedance:\n- Phase current: $I_p = I_L / \\sqrt{3} = 16.67 / \\sqrt{3} = 9.62\\text{ A}$\n- Phase voltage: $V_p = V_L = 380\\text{ V}$\n- Impedance magnitude: $|Z_\\Delta| = 380 / 9.62 = 39.5\\ \\Omega$\n- Angle: $\\theta = \\arccos(0.8) = 36.87^\\circ$\n- $Z_\\Delta = 39.5\\angle 36.87^\\circ = 31.6 + j23.7\\ \\Omega$\n($R_\\Delta = 31.6\\ \\Omega, X_{L\\Delta} = 23.7\\ \\Omega$).",
        "pitfall": "In Delta, $V_p = V_L$, but $I_p = I_L / \\sqrt{3}$. Do not divide voltage by $\\sqrt{3}$!",
        "mastered": False
    },
    {
        "id": "CH06-09",
        "chapter": "Chapter 6",
        "topic": "Two-Wattmeter Method for 3-Phase Power Measurement",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "Lecture Slides & Jele Formula Sheet",
        "question": "Explain the Two-Wattmeter Method for measuring balanced 3-phase power. State the formulas for total active power ($P$), reactive power ($Q$), and power factor angle ($\\theta$) from meter readings $W_1$ and $W_2$.",
        "formula": "$P_{\\text{total}} = W_1 + W_2$\n$Q_{\\text{total}} = \\sqrt{3}(W_1 - W_2)$\n$\\tan\\theta = \\sqrt{3}\\frac{W_1 - W_2}{W_1 + W_2}$",
        "solution": "1. Total Real Power: $P = W_1 + W_2$\n2. Total Reactive Power: $Q = \\sqrt{3}(W_1 - W_2)$\n3. Power Factor Angle: $\\theta = \\arctan\\left(\\sqrt{3}\\frac{W_1 - W_2}{W_1 + W_2}\\right)$\n4. Power Factor: $PF = \\cos\\theta$\n\nSpecial Cases:\n- If $W_1 = W_2$: $\\theta = 0^\\circ \\implies PF = 1.0$ (purely resistive).\n- If $W_2 = 0$: $PF = 0.5$ ($\\\\theta = 60^\\circ$).\n- If $W_2 < 0$: $PF < 0.5$ (one wattmeter reading turns negative).",
        "pitfall": "If $PF < 0.5$, one wattmeter deflects backward. In calculations, you MUST preserve the negative sign (e.g. $W_1 + (-W_2)$).",
        "mastered": False
    },
    {
        "id": "CH06-10",
        "chapter": "Chapter 6",
        "topic": "Neutral Wire Current in Balanced vs Unbalanced Systems",
        "level": "Level 1 (Concept)",
        "exam_reference": "HW Ch.6 Problem 1, 2, 3",
        "question": "What is the neutral current $I_N$ in a balanced 4-wire Y-Y system? What causes current to flow in the neutral wire in an unbalanced system?",
        "formula": "$I_N = -(I_a + I_b + I_c)$",
        "solution": "1. Balanced System: $I_a + I_b + I_c = 0$, therefore neutral current is strictly ZERO ($I_N = 0$). No current returns through the neutral line.\n2. Unbalanced System: The phase currents do not sum to zero due to unequal load impedances ($Z_A \\neq Z_B \\neq Z_C$), so neutral current $I_N = -(I_a + I_b + I_c) \\neq 0$ carries the return unbalance current.",
        "pitfall": "If the neutral wire is disconnected in an unbalanced load, the neutral point shifts ($V_{Nn} \\neq 0$), causing severe overvoltage across the lighter load phase!",
        "mastered": False
    },

    # =========================================================================
    # CHAPTER 7: ELECTRICAL MACHINES (DC & AC) (15 CARDS)
    # =========================================================================
    {
        "id": "CH07-01",
        "chapter": "Chapter 7",
        "topic": "Right-Hand Grip Rule for Magnetic Field around a Conductor",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.1 (Page 28 - 2 Points)",
        "question": "Current flows upward along a straight vertical wire. Using the right-hand grip rule, sketch/state the direction of the magnetic field on the left and right sides of the wire.\n" + make_img_html('exam_p28_right_hand.png', 'Exam Problem 4.1: Magnetic Field around Current Conductor'),
        "formula": "Right-Hand Grip Rule: Thumb = Current ($I$), Fingers = Field ($B$)",
        "solution": "Point right thumb upward along the current flow:\n- Left side of the wire: Magnetic field lines point OUT OF THE PAGE (represented by dots •).\n- Right side of the wire: Magnetic field lines point INTO THE PAGE (represented by crosses $\\times$).",
        "pitfall": "Direct 2-point exam question! Remember: Dot (•) = coming towards you / out of paper; Cross (×) = going away from you / into paper.",
        "mastered": False
    },
    {
        "id": "CH07-02",
        "chapter": "Chapter 7",
        "topic": "Lorentz Force on Current-Carrying Conductor in DC Motor",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.2 (Page 28 - 4 Points)",
        "question": "State the formula and rule used to determine the electromagnetic force developed on a conductor in a DC motor armature.",
        "formula": "$\\vec{F} = I(\\vec{L} \\times \\vec{B}) \\implies F = B I L \\sin\\alpha$",
        "solution": "1. Formula: $F = B I L \\sin\\alpha$, where $B$ is flux density (Tesla), $I$ is conductor current (A), $L$ is active conductor length (m), and $\\alpha$ is the angle between wire and magnetic field (usually $90^\\circ$, so $F = BIL$).\n2. Direction: Determined by Fleming's Left-Hand Rule for Motors:\n- Thumb = Motion / Force ($F$)\n- First finger = Magnetic Field ($B$, North to South)\n- Second finger = Current ($I$, positive to negative).",
        "pitfall": "Use LEFT hand for motors (converting electrical to mechanical motion); use RIGHT hand for generators (inducing EMF).",
        "mastered": False
    },
    {
        "id": "CH07-03",
        "chapter": "Chapter 7",
        "topic": "DC Machine Stator and Rotor Structural Components",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.3 (Page 28 - 4 Points)",
        "question": "What are the two main structural parts of a DC machine? Name at least 2 components found in each part.",
        "formula": "Machine = Stator (Stationary) + Rotor (Rotating)",
        "solution": "1. Stator (Stationary Part):\n- Frame / Yoke (cast iron outer housing and magnetic return path)\n- Main Field Poles & Pole Shoes (laminated sheet steel)\n- Field Windings (copper coils producing main flux $\\Phi$)\n- End covers and brush holders.\n\n2. Rotor / Armature (Rotating Part):\n- Armature Core (laminated silicon steel with slots)\n- Armature Windings (copper coils carrying load current)\n- Commutator (segmented copper bars with mica insulation)\n- Shaft, bearings, and cooling fan.",
        "pitfall": "Armature core is made of laminated silicon steel sheets specifically to minimize Eddy current losses and Hysteresis losses.",
        "mastered": False
    },
    {
        "id": "CH07-04",
        "chapter": "Chapter 7",
        "topic": "DC Machine Back EMF Formula & Physical Meaning",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides & Jele Formula Sheet",
        "question": "State the equation for induced armature back-EMF ($E_a$) in a DC machine and define all variables. How does it relate to motor speed?",
        "formula": "$E_a = K_a \\Phi \\omega_m = \\left(\\frac{p Z}{2\\pi a}\\right) \\Phi \\omega_m$",
        "solution": "- $E_a$ = Induced armature back EMF (V)\n- $K_a = \\frac{p Z}{2\\pi a}$ = Machine design constant ($p$ = poles, $Z$ = total conductors, $a$ = parallel paths)\n- $\\Phi$ = Magnetic flux per pole (Wb)\n- $\\omega_m$ = Mechanical angular velocity (rad/s) = $\\frac{2\\pi n}{60}$ ($n$ in rpm).\n\nPhysical Meaning: In a motor, $E_a$ opposes applied terminal voltage ($V_t$). As speed increases, $E_a$ increases, limiting armature current $I_a = (V_t - E_a)/R_a$.",
        "pitfall": "At starting ($n = 0$), $E_a = 0\\text{ V}$. Without starting resistors, inrush current $I_{\\text{start}} = V_t / R_a$ can reach 10-20 times rated current!",
        "mastered": False
    },
    {
        "id": "CH07-05",
        "chapter": "Chapter 7",
        "topic": "DC Machine Developed Torque & Converted Power",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides & Homework 7",
        "question": "State the formula for developed electromagnetic torque ($T_d$) in a DC motor and the relationship between developed mechanical power and electrical power.",
        "formula": "$T_d = K_a \\Phi I_a$\n$P_{\\text{developed}} = E_a I_a = T_d \\omega_m$",
        "solution": "1. Developed Torque: $T_d = K_a \\Phi I_a$ [$\\text{N}\\cdot\\text{m}$]. It is strictly proportional to field flux and armature current.\n2. Converted Power: $P_d = E_a I_a = T_d \\omega_m$ [W]. The internal electrical power converted equals mechanical power developed.\n3. Shaft Output Power: $P_{\\text{shaft}} = P_d - P_{\\text{rotational losses}}$.",
        "pitfall": "Do not confuse developed torque ($T_d$) with shaft torque ($T_{\\text{shaft}}$). Rotational losses (friction + windage) reduce output torque.",
        "mastered": False
    },
    {
        "id": "CH07-06",
        "chapter": "Chapter 7",
        "topic": "DC Shunt Motor Equivalent Circuit & KVL Equations",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 3 & Lecture Slides",
        "question": "Draw the equivalent circuit of a DC Shunt Motor and state the KVL voltage loop equation and KCL current node equation.\n" + make_img_html('dc_motor_circuit.png', 'DC Shunt Motor Equivalent Circuit & KVL'),
        "formula": "$V_t = E_a + I_a R_a$\n$I_t = I_a + I_f$\n$I_f = V_t / R_f$",
        "solution": "1. Field Circuit (Shunt across terminals):\n$I_f = \\frac{V_t}{R_f}$\n\n2. Armature Circuit (KVL):\n$$V_t = E_a + I_a R_a \\implies E_a = V_t - I_a R_a$$\n\n3. Current Relation (KCL at input terminal):\n$$I_t = I_a + I_f \\implies I_a = I_t - I_f$$\n\n4. Converted Mechanical Power:\n$$P_{\\text{conv}} = E_a I_a$$.",
        "pitfall": "In a MOTOR, $V_t > E_a$ because the source drives current into the machine ($V_t = E_a + I_a R_a$).",
        "mastered": False
    },
    {
        "id": "CH07-07",
        "chapter": "Chapter 7",
        "topic": "DC Shunt Motor Solved Example (3 hp, 120 V)",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 3 (Example 1)",
        "question": "A $3\\text{ hp}, 120\\text{ V}$ DC shunt motor has armature resistance $R_a = 0.2\\ \\Omega$, field resistance $R_f = 60\\ \\Omega$, and draws a full-load line current $I_t = 24.5\\text{ A}$. Calculate: (1) Field current $I_f$, (2) Armature current $I_a$, (3) Back EMF $E_a$, (4) Converted mechanical power $P_{\\text{conv}}$, (5) Rotational loss if output is rated at $3\\text{ hp}$. ($1\\text{ hp} = 746\\text{ W}$)",
        "formula": "$I_f = V_t / R_f$\n$I_a = I_t - I_f$\n$E_a = V_t - I_a R_a$\n$P_{\\text{conv}} = E_a I_a$\n$P_{\\text{rot}} = P_{\\text{conv}} - P_{\\text{out}}$",
        "solution": "1. Field Current:\n$$I_f = \\frac{120\\text{ V}}{60\\ \\Omega} = 2.0\\text{ A}$$\n\n2. Armature Current:\n$$I_a = I_t - I_f = 24.5 - 2.0 = 22.5\\text{ A}$$\n\n3. Back EMF:\n$$E_a = V_t - I_a R_a = 120 - (22.5 \\times 0.2) = 120 - 4.5 = 115.5\\text{ V}$$\n\n4. Converted Power:\n$$P_{\\text{conv}} = E_a I_a = 115.5\\text{ V} \\times 22.5\\text{ A} = 2,598.75\\text{ W}$$\n\n5. Rotational Loss:\n$P_{\\text{out}} = 3 \\times 746 = 2,238\\text{ W}$\n$$P_{\\text{rot}} = P_{\\text{conv}} - P_{\\text{out}} = 2,598.75 - 2,238 = 360.75\\text{ W}$$.",
        "pitfall": "Always separate internal converted power ($P_{conv} = E_a I_a$) from output shaft power ($P_{out}$). The difference is rotational friction and windage loss.",
        "mastered": False
    },
    {
        "id": "CH07-08",
        "chapter": "Chapter 7",
        "topic": "DC Shunt Motor Speed Control & Series Resistance",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 5 (Example 3a & 3b)",
        "question": "A $220\\text{ V}$ DC shunt motor has $R_a = 0.5\\ \\Omega, R_f = 110\\ \\Omega$. At full load, $I_t = 42\\text{ A}$ and speed $n = 1,200\\text{ rpm}$. (1) Find developed torque $T_d$ and constant $K_a\\Phi$. (2) If series resistor $R_{\\text{adj}} = 1.0\\ \\Omega$ is added in the armature branch and torque drops to $T_d' = 20\\text{ N}\\cdot\\text{m}$, calculate the new speed $n'$.\n" + make_img_html('dc_shunt_speed_control.png', 'DC Shunt Motor Speed Control with Added Resistance'),
        "formula": "$I_f = V_t / R_f, I_a = I_t - I_f$\n$E_a = V_t - I_a R_a$\n$\\omega_m = \\frac{2\\pi n}{60}, K_a\\Phi = E_a / \\omega_m$\n$I_a' = T_d' / (K_a\\Phi)$\n$E_a' = V_t - I_a'(R_a + R_{\\text{adj}})$\n$n' = \\frac{60 E_a'}{2\\pi(K_a\\Phi)}$",
        "solution": "1. Initial Condition:\n- $I_f = 220 / 110 = 2.0\\text{ A}$\n- $I_a = 42 - 2 = 40.0\\text{ A}$\n- $E_a = 220 - (40 \\times 0.5) = 200.0\\text{ V}$\n- $\\omega_m = \\frac{2\\pi(1200)}{60} = 125.66\\text{ rad/s}$\n- $K_a\\Phi = \\frac{E_a}{\\omega_m} = \\frac{200}{125.66} = 1.592\\text{ V}\\cdot\\text{s/rad}$\n- $T_d = K_a\\Phi I_a = 1.592 \\times 40 = 63.66\\text{ N}\\cdot\\text{m}$\n\n2. With $R_{\\text{adj}} = 1.0\\ \\Omega$ and $T_d' = 20\\text{ N}\\cdot\\text{m}$:\n- New armature current: $I_a' = \\frac{20}{1.592} = 12.56\\text{ A}$\n- Total armature resistance: $R_{\\text{tot}} = 0.5 + 1.0 = 1.5\\ \\Omega$\n- New Back EMF: $E_a' = 220 - (12.56 \\times 1.5) = 220 - 18.84 = 201.16\\text{ V}$\n- New Speed: $\\omega_m' = \\frac{201.16}{1.592} = 126.36\\text{ rad/s}$\n- In RPM: $n' = \\frac{60 \\times 126.36}{2\\pi} = 1,206.6\\text{ rpm}$.",
        "pitfall": "When $R_{adj}$ is inserted in the armature branch, total resistance becomes $R_a + R_{adj}$. Notice how reducing torque requirement allows speed to remain high even with added resistance.",
        "mastered": False
    },
    {
        "id": "CH07-09",
        "chapter": "Chapter 7",
        "topic": "DC Shunt Generator KVL, No-Load to Full-Load",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 4 (Example 2)",
        "question": "A DC shunt generator driven at $1,800\\text{ rpm}$ produces terminal voltage $V_t = 115\\text{ V}$ at no load ($I_a = 0$). When loaded, it delivers $I_a = 40\\text{ A}$ at $V_t = 100\\text{ V}$. Flux per pole is $\\Phi = 150\\text{ mWb}$. Find: (1) Machine constant $K_a$, (2) Armature resistance $R_a$, (3) Developed torque $T_d$.\n" + make_img_html('dc_generator_circuit.png', 'DC Generator Circuit and KVL'),
        "formula": "No load: $E_a = V_t = 115\\text{ V}$\n$\\omega_m = \\frac{2\\pi n}{60} = 188.5\\text{ rad/s}$\n$K_a = \\frac{E_a}{\\Phi \\omega_m}$\nLoaded: $E_a = V_t + I_a R_a \\implies R_a = \\frac{E_a - V_t}{I_a}$\n$T_d = K_a \\Phi I_a$",
        "solution": "1. Angular speed: $\\omega_m = \\frac{2\\pi \\times 1800}{60} = 188.5\\text{ rad/s}$\n2. At no load ($I_a = 0$), $E_a = V_t = 115\\text{ V}$:\n$$K_a = \\frac{115}{150 \\times 10^{-3} \\times 188.5} = \\frac{115}{28.27} = 4.07$$\n3. Loaded operation: Assuming speed and flux remain constant, $E_a = 115\\text{ V}$.\n$$R_a = \\frac{E_a - V_t}{I_a} = \\frac{115 - 100}{40} = \\frac{15}{40} = 0.375\\ \\Omega$$\n4. Developed Torque:\n$$T_d = K_a \\Phi I_a = 4.07 \\times (150 \\times 10^{-3}) \\times 40 = 24.42\\text{ N}\\cdot\\text{m}$$.",
        "pitfall": "Crucial distinction: In a GENERATOR: $E_a = V_t + I_a R_a$ ($E_a > V_t$). In a MOTOR: $E_a = V_t - I_a R_a$ ($V_t > E_a$).",
        "mastered": False
    },
    {
        "id": "CH07-10",
        "chapter": "Chapter 7",
        "topic": "DC Shunt Motor Equivalent Circuit & Speed/Torque Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.7 Problem 4 & Talui Homework Page 1",
        "question": "A 2-pole, 2.5 hp, 100 V DC shunt motor has $K_a = 153$, full-load total current $I_t = 25.0\\text{ A}$, armature resistance $R_a = 0.75\\ \\Omega$, field flux $\\Phi = 10.0\\text{ mWb}$, and field resistance $R_f = 440\\ \\Omega$. Calculate: (1) Armature current ($I_a$), (2) Back EMF ($E_a$), (3) Motor speed ($n$ in rpm), (4) Developed torque ($T_d$).",
        "formula": "$I_f = V_t / R_f$\n$I_a = I_t - I_f$\n$E_a = V_t - I_a R_a$\n$\\omega_m = E_a / (K_a \\Phi) \\implies n = \\frac{60 \\omega_m}{2\\pi}$\n$T_d = K_a \\Phi I_a$",
        "solution": "1. Field current: $I_f = 100 / 440 = 0.23\\text{ A}$\n2. Armature current: $I_a = I_t - I_f = 25.0 - 0.23 = 24.77\\text{ A}$\n3. Back EMF: $E_a = V_t - I_a R_a = 100 - (24.77 \\times 0.75) = 100 - 18.58 = 81.42\\text{ V}$\n4. Speed calculation:\n$$E_a = K_a \\Phi \\omega_m = 153 \\times (10 \\times 10^{-3}) \\times \\left(\\frac{2\\pi n}{60}\\right)$$\n$$81.42 = 1.53 \\times 0.10472 \\times n = 0.16022 \\times n \\implies n = 508.17\\text{ rpm}$$\n5. Developed Torque:\n$$T_d = K_a \\Phi I_a = 153 \\times (10 \\times 10^{-3}) \\times 24.77 = 37.9\\text{ N}\\cdot\\text{m}$$.",
        "pitfall": "Always subtract field current $I_f$ from total supply current $I_t$ before computing $I_a R_a$ drop! Field current does not flow through armature.",
        "mastered": False
    },
    {
        "id": "CH07-11",
        "chapter": "Chapter 7",
        "topic": "DC Series Motor Comparison & Runaway Hazard",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.7 Problem 5 & Talui Homework Page 2-3",
        "question": "A 2.5 hp series motor has $K_a = 153, R_a = 0.75\\ \\Omega, \\Phi = 10.0\\text{ mWb}, R_f = 440\\ \\Omega$. Find the required terminal voltages: (a) $V_{t1}$ to give the same speed ($508.17\\text{ rpm}$) as Problem 4, (b) $V_{t2}$ to give the same torque ($37.9\\text{ N}\\cdot\\text{m}$) as Problem 4.",
        "formula": "$V_t = I_a(R_a + R_f) + E_a$\n(a) Same speed: $E_a = 81.42\\text{ V}, I_a = P / E_a = 1865 / 81.42 = 22.91\\text{ A}$\n(b) Same torque: $I_a = 24.77\\text{ A}, E_a = P / I_a = 1865 / 24.77 = 75.3\\text{ V}$",
        "solution": "(a) For same speed ($n = 508.17\\text{ rpm}$):\n$E_a = 81.42\\text{ V}$\n$I_a = P_e / E_a = 1865 / 81.42 = 22.91\\text{ A}$\n$$V_{t1} = I_a(R_f + R_a) + E_a = 22.91(440 + 0.75) + 81.42 = 10,181\\text{ V} = 10.18\\text{ kV}$$\n\n(b) For same torque ($T_d = 37.9\\text{ N}\\cdot\\text{m}$):\n$I_a = 24.77\\text{ A}$\n$E_a = P_e / I_a = 1865 / 24.77 = 75.3\\text{ V}$\n$$V_{t2} = I_a(R_f + R_a) + E_a = 24.77(440.75) + 75.3 = 10,993\\text{ V} \\approx 11.0\\text{ kV}$$\n\n(c) Can same speed and same torque happen simultaneously? NO, because speed and torque uniquely dictate $E_a$ and $I_a$.",
        "pitfall": "Why must a DC series motor never run at no load? Because $I_a \\to 0 \\implies \\Phi \\to 0 \\implies \\omega_m \\approx V_t / (K_a \\Phi) \\to \\infty$ (motor dangerously overspeeds and tears itself apart).",
        "mastered": False
    },
    {
        "id": "CH07-12",
        "chapter": "Chapter 7",
        "topic": "Squirrel Cage Induction Motor & Reversing Rotation",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.4 (Page 27 - 8 Points)",
        "question": "What is a squirrel-cage rotor in a 3-phase induction motor? Explain its operating principle and describe how to reverse its direction of rotation.",
        "formula": "Swap any 2 supply lines (Phase A and Phase B)",
        "solution": "1. Structure: Rotor composed of thick copper or aluminum conductor bars embedded in a laminated steel core, permanently short-circuited at both ends by heavy conductive end-rings (resembling a squirrel exercise wheel).\n2. Operating Principle: 3-phase AC in stator windings produces a Rotating Magnetic Field (RMF). The RMF sweeps across the rotor bars, inducing an EMF and heavy circulating currents. Interaction between rotor currents and the RMF produces driving torque following Lenz's Law.\n3. Reversing Direction: Swap ANY TWO of the three incoming supply lines (e.g. swap A and B). This reverses the phase sequence from A-B-C to A-C-B, reversing the RMF rotation direction.",
        "pitfall": "Direct 8-point exam question! Never say 'reverse DC polarity' because this is an AC motor without brushes or commutators.",
        "mastered": False
    },
    {
        "id": "CH07-13",
        "chapter": "Chapter 7",
        "topic": "Synchronous Speed, Rotor Speed, and Slip",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides & Jele Formula Sheet",
        "question": "Define Synchronous Speed ($N_s$) and Slip ($s$) in a 3-phase induction motor. Why can an induction motor never run at synchronous speed under normal motoring conditions?",
        "formula": "$N_s = \\frac{120 f}{p}\\quad [\\text{rpm}]$\n$s = \\frac{N_s - N_r}{N_s} = \\frac{\\omega_s - \\omega_r}{\\omega_s}$",
        "solution": "1. Synchronous Speed: $N_s = \\frac{120 f}{p}$ is the rotational speed of the stator magnetic field ($f$ = grid frequency in Hz, $p$ = total poles).\n2. Slip: $s = \\frac{N_s - N_r}{N_s}$, typically $2\\% - 5\\%$ at rated full load.\n3. Why $N_r < N_s$: If the rotor ran at synchronous speed ($N_r = N_s$), the relative motion between rotor bars and magnetic field would become zero ($d\\Phi/dt = 0$). No EMF would be induced, zero rotor current would flow, and developed torque would drop to zero. The rotor must slip behind $N_s$ to produce torque.",
        "pitfall": "If $p$ is given as pole pairs $P$, the formula is $N_s = \\frac{60 f}{P}$. Always check whether the question specifies poles or pole pairs.",
        "mastered": False
    },
    {
        "id": "CH07-14",
        "chapter": "Chapter 7",
        "topic": "Types of Single-Phase Induction Motors",
        "level": "Level 1 (Concept)",
        "exam_reference": "AC Machines Slides P.12-14",
        "question": "Single-phase induction motors are not self-starting. Name the 3 common types used to create starting torque and describe their operating features.",
        "formula": "Shaded-Pole | Capacitor-Start | Universal Motor",
        "solution": "1. Shaded-Pole Motor:\n- Uses copper shading rings on pole faces to create a delayed phase shift in magnetic flux.\n- Low starting torque, simple, low cost; used in small desk fans and blowers.\n2. Capacitor-Start Motor:\n- Uses an auxiliary starting winding in series with a capacitor and centrifugal switch.\n- High starting torque ($200-350\\%$ of rated); switch disconnects capacitor at $\\approx 75\\%$ speed; used in air compressors, water pumps, refrigerators.\n3. Universal Motor:\n- Series-wound motor that operates on both AC and DC.\n- Very high speed (up to 20,000 rpm) and high torque-to-weight ratio; used in vacuum cleaners and power drills.",
        "pitfall": "Pure single-phase AC creates a pulsating (standing) magnetic field, not a rotating magnetic field. It requires an auxiliary phase or shading ring to start rotating.",
        "mastered": False
    },
    {
        "id": "CH07-15",
        "chapter": "Chapter 7",
        "topic": "Synchronous Motor vs Induction Motor Operation",
        "level": "Level 1 (Concept)",
        "exam_reference": "AC Machines Slides P.15 & HW7 Problem 3",
        "question": "Compare the rotor construction and operating speed of a Synchronous Motor with an Induction Motor. How does a synchronous motor maintain exact synchronous speed?",
        "formula": "$N_r = N_s$ (Synchronous) vs $N_r < N_s$ (Induction)",
        "solution": "1. Synchronous Motor:\n- Rotor is excited by an external DC source (or permanent magnets) to establish fixed magnetic poles.\n- Rotor magnetic poles magnetically lock with the stator's Rotating Magnetic Field (RMF).\n- Runs at exactly synchronous speed ($N_r = N_s = \\frac{120f}{p}$) under all steady-load conditions (zero slip, $s = 0$).\n\n2. Induction Motor:\n- Rotor currents are induced purely by relative motion between RMF and rotor conductors.\n- Must always run below synchronous speed ($N_r < N_s$) to produce torque.",
        "pitfall": "Synchronous motors can also operate at leading power factor by over-exciting the DC field, functioning as a Synchronous Condenser for power factor correction.",
        "mastered": False
    },

    # =========================================================================
    # CHAPTER 8: TRANSFORMERS (11 CARDS)
    # =========================================================================
    {
        "id": "CH08-01",
        "chapter": "Chapter 8",
        "topic": "Transformer Definition, Purpose & Structure",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 1.1 & 1.2 (Page 32 - 4 Points)",
        "question": "What is an electrical transformer? What is its primary function and structure? Draw/describe its essential construction.\n" + make_img_html('exam_p32_transformer_core.png', 'Exam Problem 1.2: Transformer Magnetic Core & Coils'),
        "formula": "Static machine transferring AC electrical energy via mutual induction.",
        "solution": "1. Definition & Purpose: A static electrical machine that transfers AC electrical power between circuits at the same frequency, stepping voltage up or down to minimize transmission loss or match load requirements.\n2. Structure: Consists of:\n- High-permeability laminated silicon steel core (provides closed magnetic circuit)\n- Primary winding ($N_1$ turns) connected to input AC source\n- Secondary winding ($N_2$ turns) connected to output load.",
        "pitfall": "A transformer changes voltage and current levels, but NEVER changes power frequency ($f_1 = f_2 = 50\\text{ Hz}$).",
        "mastered": False
    },
    {
        "id": "CH08-02",
        "chapter": "Chapter 8",
        "topic": "Working Principle & Induced EMF Equation",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 1.3 (Page 32 - 4 Points)",
        "question": "Explain the working principle of a transformer and write Faraday's induced EMF formula including the derivation constant.",
        "formula": "$E_{\\text{rms}} = 4.44 f N \\Phi_m = \\sqrt{2}\\pi f N \\Phi_m$",
        "solution": "1. Working Principle: Mutual electromagnetic induction according to Faraday's law. When AC voltage $V_1$ is applied to the primary winding, it drives alternating current that establishes a time-varying magnetic flux $\\Phi(t) = \\Phi_m \\sin(\\omega t)$ in the core. This flux links both windings, inducing EMF $e_1 = -N_1 \\frac{d\\Phi}{dt}$ and $e_2 = -N_2 \\frac{d\\Phi}{dt}$.\n2. RMS EMF Equation: Since $\\frac{d}{dt}[\\sin(\\omega t)] = \\omega \\cos(\\omega t)$, peak value is $E_m = 2\\pi f N \\Phi_m$. Dividing by $\\sqrt{2}$ yields:\n$$E_{\\text{rms}} = \\frac{2\\pi}{\\sqrt{2}} f N \\Phi_m = \\sqrt{2}\\pi f N \\Phi_m = 4.44 f N \\Phi_m$$.",
        "pitfall": "Remember the constant 4.44 comes from $\\sqrt{2}\\pi \\approx 1.4142 \\times 3.1416 = 4.4428$.",
        "mastered": False
    },
    {
        "id": "CH08-03",
        "chapter": "Chapter 8",
        "topic": "Ideal Transformer 4 Essential Properties",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 1.4 (Page 32 - 2 Points)",
        "question": "List the four fundamental assumptions/properties of an Ideal Transformer.",
        "formula": "1. Zero resistance, 2. Zero leakage, 3. Zero core loss, 4. Infinite permeability",
        "solution": "1. Zero winding resistance: $R_1 = 0, R_2 = 0$ (no $I^2R$ copper losses; $V_1 = E_1, V_2 = E_2$).\n2. Zero leakage flux: All magnetic flux is completely confined within the core and links 100% of both primary and secondary windings ($L_1 = 0, L_2 = 0$).\n3. Zero core loss: No hysteresis loss and no eddy current loss in the core ($P_{\\text{core}} = 0$, efficiency $\\eta = 100\\%$).\n4. Infinite core permeability: $\\mu_r \\to \\infty$, meaning zero magnetizing current is required to establish core flux.",
        "pitfall": "Direct 2-point exam question! Make sure to list all 4 properties clearly.",
        "mastered": False
    },
    {
        "id": "CH08-04",
        "chapter": "Chapter 8",
        "topic": "Power Losses in a Practical Transformer",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 1.5 (Page 32 - 4 Points)",
        "question": "What types of power losses occur in a real transformer? Name the 2 primary loss categories and where they occur.",
        "formula": "$P_{\\text{loss}} = P_{\\text{copper}} + P_{\\text{core}} = (I_1^2 R_1 + I_2^2 R_2) + (P_h + P_e)$",
        "solution": "1. Copper Losses (Winding Losses): Occur in the copper windings due to Joule heating: $P_{\\text{cu}} = I_1^2 R_1 + I_2^2 R_2$. Variable loss proportional to $(I_{\\text{load}})^2$.\n2. Core / Iron Losses: Occur in the magnetic core, nearly constant at all load levels:\n- Hysteresis Loss: Energy lost in continuously reversing magnetic domains ($P_h \\propto f B_m^{1.6}$).\n- Eddy Current Loss: Circulating currents induced in the conductive iron core ($P_e \\propto f^2 B_m^2 t^2$). Minimized by using thin laminated sheets insulated by varnish.",
        "pitfall": "Core loss depends on voltage and frequency (constant 24/7); copper loss depends on load current squared.",
        "mastered": False
    },
    {
        "id": "CH08-05",
        "chapter": "Chapter 8",
        "topic": "Step-Down Transformer Parameter Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "Final Exam Problem 1.6 (Page 31 - 12 Points)",
        "question": "A single-phase transformer rated at $25\\text{ kVA}, 50\\text{ Hz}$ has primary winding $N_1 = 500\\text{ turns}$ and secondary winding $N_2 = 50\\text{ turns}$. The primary winding is connected to $V_1 = 3,000\\text{ V}$. Calculate: (a) Turns ratio ($a$), (b) Full-load rated primary and secondary currents ($I_1, I_2$), (c) Induced secondary voltage ($V_2$), (d) Maximum magnetic flux in core ($\\Phi_m$).",
        "formula": "$a = N_1 / N_2$\n$I_1 = S / V_1, I_2 = S / V_2$\n$V_2 = V_1 / a$\n$\\Phi_m = \\frac{V_1}{4.44 f N_1}$",
        "solution": "(a) Turns Ratio:\n$$a = \\frac{N_1}{N_2} = \\frac{500}{50} = 10$$\n\n(b) Rated Full-Load Currents:\n$$I_1 = \\frac{S_{\\text{rated}}}{V_1} = \\frac{25,000}{3,000} = 8.33\\text{ A}$$\n$$I_2 = a \\times I_1 = 10 \\times 8.33 = 83.33\\text{ A}$$\n\n(c) Secondary Voltage:\n$$V_2 = \\frac{V_1}{a} = \\frac{3,000}{10} = 300\\text{ V}$$\n\n(d) Maximum Magnetic Flux:\n$$\\Phi_m = \\frac{V_1}{4.44 f N_1} = \\frac{3,000}{4.44 \\times 50 \\times 500} = \\frac{3,000}{111,000} = 0.0270\\text{ Wb} = 27.0\\text{ mWb}$$.",
        "pitfall": "Exact problem from past exam page 31 (12 points total)! Write all steps and units clearly ($A, V, Wb$).",
        "mastered": False
    },
    {
        "id": "CH08-06",
        "chapter": "Chapter 8",
        "topic": "140V/80V Transformer Turns & Currents",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.8 Problem 2 & Talui Homework Page 4",
        "question": "A single-phase transformer has primary and secondary voltages $V_1 = 140\\text{ V}$ and $V_2 = 80\\text{ V}$. Primary current is $I_1 = 8\\text{ A}$. The secondary winding has $N_2 = 800\\text{ turns}$. Find: (1) Primary winding turns $N_1$, (2) Secondary current $I_2$, (3) Turns ratio $a$.",
        "formula": "$a = V_1 / V_2 = N_1 / N_2$\n$I_2 = a \\times I_1$",
        "solution": "1. Turns Ratio:\n$$a = \\frac{V_1}{V_2} = \\frac{140}{80} = 1.75$$\n\n2. Primary Winding Turns:\n$$N_1 = a \\times N_2 = 1.75 \\times 800 = 1,400\\text{ turns}$$\n\n3. Secondary Winding Current:\n$$I_2 = a \\times I_1 = 1.75 \\times 8\\text{ A} = 14.0\\text{ A}$$.",
        "pitfall": "Notice that voltage drops by 1.75 times, so current increases by 1.75 times to conserve total volt-amperes ($S = V_1 I_1 = V_2 I_2 = 1,120\\text{ VA}$).",
        "mastered": False
    },
    {
        "id": "CH08-07",
        "chapter": "Chapter 8",
        "topic": "200 kVA, 2200V/220V Transformer Full Calculations",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.8 Problem 3 & Talui Homework Page 4",
        "question": "A single-phase transformer rated at $200\\text{ kVA}, 2,200\\text{ V} / 220\\text{ V}, 50\\text{ Hz}$ has primary winding $N_1 = 200\\text{ turns}$. Find: (1) Rated primary and secondary currents $I_1, I_2$, (2) Secondary winding turns $N_2$, (3) Maximum magnetic flux in core $\\Phi_m$.",
        "formula": "$I_1 = S / V_1, I_2 = S / V_2$\n$a = V_1 / V_2 = N_1 / N_2$\n$\\Phi_m = \\frac{V_1}{4.44 f N_1}$",
        "solution": "1. Rated Currents:\n$$I_1 = \\frac{200,000\\text{ VA}}{2,200\\text{ V}} = 90.91\\text{ A}$$\n$$I_2 = \\frac{200,000\\text{ VA}}{220\\text{ V}} = 909.09\\text{ A}$$\n\n2. Secondary Turns:\nTurns ratio $a = \\frac{2,200}{220} = 10$\n$$N_2 = \\frac{N_1}{a} = \\frac{200}{10} = 20\\text{ turns}$$\n\n3. Maximum Core Flux:\n$$\\Phi_m = \\frac{V_1}{4.44 f N_1} = \\frac{2,200}{4.44 \\times 50 \\times 200} = \\frac{2,200}{44,400} = 0.0495\\text{ Wb} = 49.5\\text{ mWb}$$.",
        "pitfall": "Always use rated apparent power in Volt-Amperes ($200\\text{ kVA} = 200,000\\text{ VA}$), not kilo-units.",
        "mastered": False
    },
    {
        "id": "CH08-08",
        "chapter": "Chapter 8",
        "topic": "Impedance Matching with Transformer for Maximum Power",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "Ch07Transformer.pdf Example 1 (Page 18) & FunE P.7",
        "question": "A $10\\text{ V}$ audio source with internal source resistance $R_s = 1\\ \\Omega$ is connected to a speaker load $R_L = 9\\ \\Omega$. (1) Find load power without transformer. (2) If an ideal transformer with turns ratio $a = 1/3$ is inserted between source and speaker, find the reflected load resistance and the new load power.\n" + make_img_html('transformer_audio_matching.png', 'Audio Source Impedance Matching via 1:3 Transformer'),
        "formula": "$R_L' = a^2 R_L = \\left(\\frac{N_1}{N_2}\\right)^2 R_L$\n$P = I^2 R$",
        "solution": "1. Without Transformer:\n$$I = \\frac{V_s}{R_s + R_L} = \\frac{10}{1 + 9} = 1.0\\text{ A}$$\n$$P_{\\text{speaker}} = I^2 R_L = (1.0)^2 \\times 9 = 9.0\\text{ W}$$\n\n2. With Transformer ($a = 1/3$):\nReflected load resistance seen by source:\n$$R_L' = a^2 R_L = \\left(\\frac{1}{3}\\right)^2 \\times 9 = \\frac{1}{9} \\times 9 = 1.0\\ \\Omega$$\nNotice $R_L' = R_s = 1.0\\ \\Omega$ (perfect Maximum Power Transfer match!).\nNew primary current:\n$$I_1 = \\frac{V_s}{R_s + R_L'} = \\frac{10}{1 + 1} = 5.0\\text{ A}$$\nPower delivered to reflected load (and thus speaker, since transformer is ideal):\n$$P_{\\text{speaker}} = I_1^2 R_L' = (5.0)^2 \\times 1.0 = 25.0\\text{ W}$$\nPower increased dramatically from $9\\text{ W}$ to $25\\text{ W}$!",
        "pitfall": "Reflected impedance scales as $a^2$: $Z_1 = a^2 Z_2$. In audio systems, transformers match unequal impedances to maximize transferred power.",
        "mastered": False
    },
    {
        "id": "CH08-09",
        "chapter": "Chapter 8",
        "topic": "DC Voltage Connection to Transformer Primary Hazard",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.8 Problem 4 & Talui Homework Page 5",
        "question": "What happens if a DC voltage (e.g. 200 V DC) is connected to the primary winding of a transformer? What will the secondary voltage and current be?",
        "formula": "$v_2 = -N_2 \\frac{d\\Phi}{dt} = 0$, $I_1 = \\frac{V_{\\text{dc}}}{R_1}$ (catastrophic overcurrent)",
        "solution": "1. Secondary Voltage: $V_2 = 0\\text{ V}$. Because DC voltage creates static, time-invariant magnetic flux ($\\frac{d\\Phi}{dt} = 0$), no voltage is induced in the secondary winding.\n2. Primary Hazard: On DC, inductive reactance is zero ($X_L = 2\\pi f L = 0$ since $f = 0$). There is no counter-EMF to oppose the source. Primary current is limited solely by tiny copper wire resistance $R_1$: $I_1 = V_{\\text{dc}} / R_1 \\gg I_{\\text{rated}}$. The primary winding rapidly overheats and burns out.",
        "pitfall": "Exam favorite concept! Always state: 1) Secondary output is zero, and 2) Primary winding burns out.",
        "mastered": False
    },
    {
        "id": "CH08-10",
        "chapter": "Chapter 8",
        "topic": "Practical Transformer Equivalent Circuit & Efficiency",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 10 (Example 5)",
        "question": "A $10\\text{ kVA}, 2200/220\\text{ V}, 50\\text{ Hz}$ practical transformer has primary parameters $R_1 = 3.5\\ \\Omega, X_1 = 6.2\\ \\Omega$ and secondary parameters $R_2 = 0.035\\ \\Omega, X_2 = 0.062\\ \\Omega$. Core loss is $P_{\\text{core}} = 150\\text{ W}$. Find equivalent impedance referred to primary, full-load copper loss, and efficiency at $PF = 1.0$ and $PF = 0.8$ lagging.\n" + make_img_html('transformer_equivalent_circuit.png', 'Practical Transformer Equivalent Circuit & KVL'),
        "formula": "$a = V_1 / V_2 = 10$\n$R_{eq1} = R_1 + a^2 R_2, X_{eq1} = X_1 + a^2 X_2$\n$I_1 = S / V_1$\n$P_{\\text{cu}} = I_1^2 R_{eq1}$\n$\\eta = \\frac{P_{\\text{out}}}{P_{\\text{out}} + P_{\\text{cu}} + P_{\\text{core}}} \\times 100\\%$",
        "solution": "1. Equivalent Parameters referred to Primary:\n$a = 2200 / 220 = 10$\n- $R_{eq1} = 3.5 + (10)^2(0.035) = 3.5 + 3.5 = 7.0\\ \\Omega$\n- $X_{eq1} = 6.2 + (10)^2(0.062) = 6.2 + 6.2 = 12.4\\ \\Omega$\n- $Z_{eq1} = 7.0 + j12.4\\ \\Omega$\n\n2. Full-Load Primary Current & Copper Loss:\n$I_1 = 10,000 / 2,200 = 4.545\\text{ A}$\n$$P_{\\text{cu}} = (4.545)^2 \\times 7.0 = 144.6\\text{ W}$$\n\n3. Efficiency at Unity PF ($PF = 1.0$):\n$P_{\\text{out}} = 10,000 \\times 1.0 = 10,000\\text{ W}$\n$$\\eta = \\frac{10,000}{10,000 + 144.6 + 150} \\times 100\\% = \\frac{10,000}{10,294.6} \\times 100\\% = 97.14\\%$$\n\n4. Efficiency at $PF = 0.8$ lagging:\n$P_{\\text{out}} = 10,000 \\times 0.8 = 8,000\\text{ W}$\n$$\\eta = \\frac{8,000}{8,000 + 144.6 + 150} \\times 100\\% = \\frac{8,000}{8,294.6} \\times 100\\% = 96.45\\%$$.",
        "pitfall": "Notice how lower power factor degrades transformer efficiency because the real power output drops while $I^2R$ copper loss remains at full rating.",
        "mastered": False
    },
    {
        "id": "CH08-11",
        "chapter": "Chapter 8",
        "topic": "Energy Conservation in Transformer Operations",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 11 & Lecture Slides",
        "question": "State three practical methods to save electrical energy in industrial transformer substations.",
        "formula": "1. Reduce no-load loss | 2. Consolidate light loads | 3. Power Factor correction",
        "solution": "1. De-energize Unloaded Transformers: Core losses occur 24 hours a day regardless of load. Switching off backup or redundant transformers during weekends/nights eliminates standing core loss.\n2. Load Consolidation: Operating two transformers at 20% load is far less efficient than shutting one down and running a single unit at 40-50% load near its peak efficiency point.\n3. Power Factor Correction: Installing shunt capacitor banks on the secondary side cancels inductive VARs, reducing secondary line current, which cuts $I^2 R$ winding copper losses and lowers transformer operating temperature.",
        "pitfall": "Peak efficiency occurs when copper loss equals core loss ($P_{cu} = P_{core}$), which is typically engineered around 50% to 75% rated load.",
        "mastered": False
    },

    # =========================================================================
    # CHAPTER 9: POWER TRANSMISSION SYSTEMS (10 CARDS)
    # =========================================================================
    {
        "id": "CH09-01",
        "chapter": "Chapter 9",
        "topic": "Structure of Electrical Power System & Voltage Levels",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 5.3 (Page 25 - 4 Points)",
        "question": "State the four stages of a power system and classify these standard Thai grid voltages: 500 kV, 22 kV, 380 V, and residential household voltage.\n" + make_img_html('power_system_stages.png', '4 Stages of Electric Power Transmission System'),
        "formula": "Generation -> Transmission -> Distribution -> Utilization",
        "solution": "1. Generation: 10 - 20 kV (alternator terminals)\n2. Transmission: 115 kV, 230 kV, 500 kV\n3. Distribution: 22 kV, 24 kV, 33 kV\n4. Utilization: 380/400 V ($3\\phi$) and 220/230 V ($1\\phi$)\n\nClassifications for Exam Problem 5.3:\n- 500 kV: Extra High Voltage (EHV)\n- 22 kV: Medium Voltage (MV)\n- 380 V: Low / Utilization Voltage\n- Residential Household Voltage: 220 V (or modern 230 V).",
        "pitfall": "Why transmit at Extra High Voltage (500 kV)? Because $P = \\sqrt{3}V_L I_L \\cos\\theta$. Stepping up $V$ reduces line current $I$ proportionally, minimizing $I^2 R$ transmission wire loss by $(1/a)^2$.",
        "mastered": False
    },
    {
        "id": "CH09-02",
        "chapter": "Chapter 9",
        "topic": "Three Electricity Authorities in Thailand & Roles",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides & Jele Formula Sheet",
        "question": "Name the three state electricity authorities in Thailand and state their specific geographic responsibilities.",
        "formula": "EGAT (Generation/Transmission) | MEA (BKK+2) | PEA (74 Provinces)",
        "solution": "1. EGAT (Electricity Generating Authority of Thailand / กฟผ.): Responsible for power generation and nationwide high-voltage transmission grid (115-500 kV); sells wholesale power to MEA, PEA, and industrial direct clients.\n2. MEA (Metropolitan Electricity Authority / กฟน.): Distributes electricity in 3 central provinces: Bangkok, Nonthaburi, and Samut Prakan.\n3. PEA (Provincial Electricity Authority / กฟภ.): Distributes electricity across the remaining 74 provinces in Thailand.",
        "pitfall": "Guaranteed exam question! Remember MEA covers strictly Bangkok, Nonthaburi, and Samut Prakan.",
        "mastered": False
    },
    {
        "id": "CH09-03",
        "chapter": "Chapter 9",
        "topic": "Fuel Mix for Power Generation in Thailand",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 5.1 (Page 25 - 2 Points)",
        "question": "What is the most widely used fuel source for electricity generation in Thailand, and what fuel source is used least?",
        "formula": "Most: Natural Gas (~55-60%) | Least: Fuel Oil / Heavy Fuel Oil",
        "solution": "- Most used: Natural Gas (ก๊าซธรรมชาติ) accounts for over 55-60% of total electrical energy in Thailand (Gulf of Thailand, Myanmar, and imported LNG).\n- Followed by: Coal / Lignite (Mae Moh power plant) and Renewable/Hydro imports from Lao PDR.\n- Least used: Heavy fuel oil (น้ำมันเตา) and diesel, due to extremely high operating costs and environmental impact.",
        "pitfall": "Do not answer solar or oil as the main source; it is overwhelmingly Natural Gas.",
        "mastered": False
    },
    {
        "id": "CH09-04",
        "chapter": "Chapter 9",
        "topic": "Base Load vs Peak Load Power Plants & Startup Times",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 5.2 (Page 25 - 4 Points)",
        "question": "Compare Base Load and Peak Load power plants. Which plant type has the shortest startup time, and which has the longest?",
        "formula": "Base Load: Steam/Coal (6-8 hrs) | Peak Load: Hydro (<5 mins)",
        "solution": "1. Base Load Plants: Run continuously 24/7 to supply minimum baseline grid demand. High capital cost, low running cost.\n- Longest Startup Time: Steam thermal power plants (e.g. Mae Moh coal/lignite) take 6 - 8 hours to bring boilers up to operating pressure from cold start.\n\n2. Peak Load Plants: Dispatched quickly to handle rapid spikes in daily demand.\n- Shortest Startup Time: Hydroelectric power plants (โรงไฟฟ้าพลังน้ำ) can synchronize and deliver full load in less than 5 minutes. Gas turbines take 10 - 15 minutes.",
        "pitfall": "Direct 4-part exam question: Shortest = Hydro; Longest = Steam; Base = Steam/Coal; Peak = Hydro/Gas turbine.",
        "mastered": False
    },
    {
        "id": "CH09-05",
        "chapter": "Chapter 9",
        "topic": "Transmission Tower Classifications & Structures",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 14 & Lecture Slides",
        "question": "Describe the structural features and conductor configurations used on Thai grid transmission towers at 500 kV, 230 kV, and 115 kV.\n" + make_img_html('transmission_towers.png', 'Transmission Tower Classifications (500 kV, 230 kV, 115 kV, 69 kV)'),
        "formula": "500 kV (Quadruple bundle) | 230 kV (Twin bundle) | 115 kV (Single conductor)",
        "solution": "1. 500 kV Towers: Massive steel lattice towers. Use Quadruple Bundle Conductors (4 sub-conductors spaced per phase) to expand effective conductor diameter, suppressing Corona discharge and Audible noise, with ~28-35 insulator discs per string.\n2. 230 kV Towers: Double-circuit lattice towers using Twin Bundle Conductors (2 sub-conductors per phase) with ~14-17 insulator discs.\n3. 115 kV Towers: Single-conductor lattice or spun concrete poles with ~7-9 insulator discs, linking sub-transmission substations to city borders.",
        "pitfall": "Bundle conductors are used at 500 kV and 230 kV primarily to reduce high-voltage Corona discharge loss and radio interference.",
        "mastered": False
    },
    {
        "id": "CH09-06",
        "chapter": "Chapter 9",
        "topic": "Customer Categories & Supply Voltage Thresholds",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 13 & PEA/MEA Regulations",
        "question": "How are electricity consumers categorized by supply voltage in Thailand? State the voltage levels supplied to: (1) Residential households, (2) Medium business (30-999 kVA), (3) Large industrial (>10,000 kVA).",
        "formula": "<30 kVA: LV 230/400 V | 30-999 kVA: MV 22/33 kV | >10 MVA: HV 115 kV",
        "solution": "1. Low Voltage Customers (< 30 kVA):\n- Residential households and small shops.\n- Supplied at $230\\text{ V}$ ($1\\phi, 2\\text{-wire}$) or $400/230\\text{ V}$ ($3\\phi, 4\\text{-wire}$).\n\n2. Medium Voltage Customers (30 - 999 kVA):\n- Commercial buildings, hotels, small factories.\n- Supplied at $22\\text{ kV}$ or $33\\text{ kV}$ (PEA) / $12\\text{ kV}$ or $24\\text{ kV}$ (MEA) with customer-owned step-down transformer.\n\n3. High Voltage Customers (> 10,000 kVA / 10 MVA):\n- Heavy chemical plants, steel mills, large industrial estates.\n- Supplied directly from $115\\text{ kV}$ or $230\\text{ kV}$ grid via private dedicated substation.",
        "pitfall": "Large customers buy power at high voltage because tariff rates per unit (kWh) are cheaper, and transmission loss over long industrial site feeders is minimized.",
        "mastered": False
    },
    {
        "id": "CH09-07",
        "chapter": "Chapter 9",
        "topic": "Power Plant Specifics in Thailand (Bhumibol, Ubolratana, Mae Moh, Nam Pong)",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunEE Final Note P.5 & Lecture Slides",
        "question": "Match the following major Thai power plants to their generation type, location, and key features: (1) Bhumibol Dam, (2) Ubolratana Dam, (3) Mae Moh, (4) Nam Pong.",
        "formula": "Bhumibol: Ping River / Tak | Ubolratana: Khon Kaen | Mae Moh: Lignite / Lampang | Nam Pong: Combined Cycle / Khon Kaen",
        "solution": "1. Bhumibol Dam (เขื่อนภูมิพล):\n- Location: Tak Province, Ping River.\n- Thailand's first multipurpose concrete arch dam; major hydro peak-load plant.\n\n2. Ubolratana Dam (เขื่อนอุบลรัตน์):\n- Location: Khon Kaen Province, Pong River.\n- First hydroelectric dam in Northeastern Thailand (Isan).\n\n3. Mae Moh Power Plant (โรงไฟฟ้าแม่เมาะ):\n- Location: Lampang Province.\n- Thailand's largest coal/lignite-fired steam thermal base-load power plant ($2,625\\text{ MW}$).\n\n4. Nam Pong Power Plant (โรงไฟฟ้าน้ำพอง):\n- Location: Khon Kaen Province ($32\\text{ km}$ from downtown KKU).\n- Combined-cycle power plant fueled by domestic natural gas from the Sinphuhorm/Nam Phong field.",
        "pitfall": "Very relevant for KKU engineering students! Ubolratana Dam and Nam Pong power plant are located directly in Khon Kaen.",
        "mastered": False
    },
    {
        "id": "CH09-08",
        "chapter": "Chapter 9",
        "topic": "Residential Solar PV Rooftop System Components",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunEE Final Note P.6-7 & Lecture Slides",
        "question": "Draw the flow path and name the 5 essential components of a grid-connected residential Solar PV system.",
        "formula": "PV Modules -> DC/DC Charge Controller -> Battery -> DC/AC Inverter -> Fusebox / Bi-directional Meter",
        "solution": "1. PV Solar Modules (Panels): Convert solar irradiance into DC electrical power.\n2. DC/DC Converter / Charge Controller: Regulates charging voltage and extracts maximum power (MPPT).\n3. Battery Storage: Stores electrical energy in DC form for night/emergency backup.\n4. DC/AC Power Inverter: Converts direct current (DC) into synchronized $230\\text{ V}, 50\\text{ Hz}$ alternating current (AC).\n5. Consumer Unit (Fusebox) & Bi-directional Net Meter: Distributes power to household loads and exports excess power to the MEA/PEA grid.",
        "pitfall": "Solar cells produce pure DC; household appliances run on AC. The Inverter is the crucial bridging component.",
        "mastered": False
    },
    {
        "id": "CH09-09",
        "chapter": "Chapter 9",
        "topic": "Energy Conservation in Household Washing Machines",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 5.5 (Page 24 - 2 Points)",
        "question": "State two practical, effective methods to conserve electrical energy when operating a household washing machine.",
        "formula": "1. Wash full rated loads | 2. Cold water wash | 3. Air dry clothes",
        "solution": "1. Wash with Full Rated Load: Consolidate dirty laundry to run at maximum rated drum capacity rather than running multiple partial loads (each cycle consumes roughly the same mechanical and pumping energy).\n2. Use Cold Water Wash: Avoid electrical water heating cycles. In washing machines with built-in heaters, heating water accounts for 80% to 90% of total electrical energy per wash cycle.\n3. Spin and Air Dry: Utilize high-speed mechanical spin extraction and sun/air drying instead of electric heated tumble dryers.",
        "pitfall": "Direct 2-point exam question from page 24!",
        "mastered": False
    },
    {
        "id": "CH09-10",
        "chapter": "Chapter 9",
        "topic": "Energy Conservation in Industrial Manufacturing Plants",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 5.6 (Page 24 - 3 Points)",
        "question": "State three practical engineering methods to reduce electrical energy consumption and costs in an industrial manufacturing plant.",
        "formula": "1. Load scheduling (TOU) | 2. Capacitor banks (PF) | 3. Transformer management",
        "solution": "1. Peak Load Shifting (TOU Tariff Optimization): Schedule energy-intensive machinery (electric arc furnaces, heavy crushers, pumping) during Off-Peak tariff hours (22:00 - 09:00 and weekends), dramatically lowering demand charges.\n2. Power Factor Correction with Capacitor Banks: Maintain overall plant power factor above $0.95$ lagging to eliminate utility reactive power penalty fees and reduce line current $I^2 R$ heat dissipation in plant distribution feeders.\n3. Transformer Management & High-Efficiency Motors: Consolidate lightly loaded transformers during slow shifts to eliminate no-load core losses, and upgrade old motors to Premium Efficiency (IE3/IE4) with Variable Speed Drives (VFDs).",
        "pitfall": "Direct 3-point exam question from page 24!",
        "mastered": False
    },

    # =========================================================================
    # CHAPTER 10: MEASURING INSTRUMENTS (14 CARDS)
    # =========================================================================
    {
        "id": "CH10-01",
        "chapter": "Chapter 10",
        "topic": "PMMC / Galvanometer Working Principle & Deflection Torque",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 3.1 (Page 30 - 4 Points)",
        "question": "Explain the working principle of a PMMC (Permanent Magnet Moving Coil) instrument. State the deflecting torque formula and deflection angle relation.",
        "formula": "$T_d = B A N I = B I L d N$\n$T_c = K \\theta \\implies \\theta = \\frac{B A N}{K} I$",
        "solution": "1. Principle: Motor principle (Lorentz force). When current $I$ flows through the moving coil situated in the permanent magnetic field, a deflecting torque $T_d$ is exerted on the coil sides.\n2. Deflecting Torque: $T_d = B A N I$, where $B$ = magnetic flux density (T), $A$ = coil area ($m^2$), $N$ = turns, $I$ = current (A).\n3. Controlling Torque: Opposed by hair springs: $T_c = K \\theta$.\n4. Steady-State Deflection: At balance ($T_d = T_c$):\n$$\\theta = \\left(\\frac{B A N}{K}\\right) I$$\nDeflection $\\theta$ is linearly proportional to current $I$, giving PMMC a perfectly UNIFORM (linear) scale.",
        "pitfall": "PMMC responds strictly to DC average current. If pure sinusoidal AC is applied, average torque is zero and pointer merely vibrates around zero.",
        "mastered": False
    },
    {
        "id": "CH10-02",
        "chapter": "Chapter 10",
        "topic": "PMMC Torque & Deflection Calculation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.10 Problem 1 & Talui Homework Page 6",
        "question": "A moving-coil instrument has $N = 150\\text{ turns}$, coil area $A = 100\\text{ mm}^2$, air gap flux density $B = 0.15\\text{ T}$, and current $I = 5\\text{ mA}$. Spring constant $K = 2 \\times 10^{-6}\\text{ N}\\cdot\\text{m/deg}$. Calculate: (1) Deflecting torque ($T_d$), (2) Deflection angle ($\\theta$ in degrees and radians).",
        "formula": "$T_d = N B A I$\n$\\theta = T_d / K$",
        "solution": "1. Deflecting Torque:\n$A = 100\\text{ mm}^2 = 100 \\times 10^{-6}\\text{ m}^2 = 10^{-4}\\text{ m}^2$\n$$T_d = 150 \\times 0.15 \\times (100 \\times 10^{-6}) \\times (5 \\times 10^{-3}) = 1.125 \\times 10^{-5}\\text{ N}\\cdot\\text{m}$$\n\n2. Deflection Angle:\n$$\\theta = \\frac{T_d}{K} = \\frac{1.125 \\times 10^{-5}}{2 \\times 10^{-6}} = 5.625^\\circ$$\nIn radians:\n$$\\theta = 5.625 \\times \\frac{\\pi}{180} = 0.0982\\text{ rad}$$.",
        "pitfall": "Always convert coil area from $\\text{mm}^2$ to $\\text{m}^2$ by multiplying by $10^{-6}$! Never forget this conversion.",
        "mastered": False
    },
    {
        "id": "CH10-03",
        "chapter": "Chapter 10",
        "topic": "DC Ammeter Design with and without CDRX",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.10 Problem 2 & Talui Homework Page 6",
        "question": "Design a DC ammeter with full scale $I = 100\\text{ mA}$ using a galvanometer with $I_g = 10\\text{ mA}$ and $R_g = 50\\ \\Omega$. Calculate shunt resistor $R_s$: (a) Without damping resistor, (b) With critical damping resistor CDRX = $400\\ \\Omega$ connected in series with coil.\n" + make_img_html('ammeter_voltmeter_basic.png', 'DC Ammeter with Shunt and CDRX Wiring'),
        "formula": "(a) $R_s = \\frac{I_g R_g}{I - I_g}$\n(b) $R_s = \\frac{I_g (R_g + \\text{CDRX})}{I - I_g}$",
        "solution": "(a) Without CDRX:\n$$R_s = \\frac{10\\text{ mA} \\times 50\\ \\Omega}{100\\text{ mA} - 10\\text{ mA}} = \\frac{0.5\\text{ V}}{0.09\\text{ A}} = 5.56\\ \\Omega$$\n\n(b) With CDRX = $400\\ \\Omega$ (connected in series with $R_g$):\nTotal meter branch resistance $R_{\\text{branch}} = R_g + \\text{CDRX} = 50 + 400 = 450\\ \\Omega$\n$$R_s = \\frac{I_g(R_g + \\text{CDRX})}{I - I_g} = \\frac{10\\text{ mA} \\times 450\\ \\Omega}{90\\text{ mA}} = \\frac{4.5\\text{ V}}{0.09\\text{ A}} = 50.0\\ \\Omega$$.",
        "pitfall": "In a DC ammeter, CDRX is connected in SERIES with the moving coil inside the meter branch; the shunt $R_s$ is in parallel with that entire branch.",
        "mastered": False
    },
    {
        "id": "CH10-04",
        "chapter": "Chapter 10",
        "topic": "DC Voltmeter Design with and without CDRX",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "HW Ch.10 Problem 3 & Talui Homework Page 7",
        "question": "Design a DC voltmeter with full scale $V = 50\\text{ V}$ using a galvanometer with $I_g = 10\\text{ mA}$ and $R_g = 50\\ \\Omega$. Calculate series multiplier $R_s$: (a) Without damping resistor, (b) With critical damping resistor CDRX = $400\\ \\Omega$ connected in parallel across the coil.\n" + make_img_html('ammeter_voltmeter_basic.png', 'DC Voltmeter with Series Multiplier and Parallel CDRX'),
        "formula": "(a) $R_s = \\frac{V}{I_g} - R_g$\n(b) $R_s = \\frac{V - I_g R_g}{I_g + I_{\\text{cdrx}}}$",
        "solution": "(a) Without CDRX:\n$$R_s = \\frac{V}{I_g} - R_g = \\frac{50}{0.010} - 50 = 5,000 - 50 = 4,950\\ \\Omega$$\n\n(b) With CDRX = $400\\ \\Omega$ (connected in PARALLEL across the coil):\n- Voltage across coil at FSD: $V_g = I_g R_g = 10\\text{ mA} \\times 50\\ \\Omega = 0.5\\text{ V}$\n- Current through CDRX: $I_{\\text{cdrx}} = V_g / \\text{CDRX} = 0.5 / 400 = 1.25\\text{ mA}$\n- Total circuit current: $I_{\\text{tot}} = I_g + I_{\\text{cdrx}} = 10\\text{ mA} + 1.25\\text{ mA} = 11.25\\text{ mA}$\n- Required series multiplier:\n$$R_s = \\frac{V - V_g}{I_{\\text{tot}}} = \\frac{50 - 0.5}{0.01125} = \\frac{49.5}{0.01125} = 4,400\\ \\Omega$$.",
        "pitfall": "In a voltmeter, CDRX is placed in PARALLEL with the coil to damp oscillation, so the series resistor must supply both coil current and CDRX current!",
        "mastered": False
    },
    {
        "id": "CH10-05",
        "chapter": "Chapter 10",
        "topic": "Multi-Range DC Ammeter Design with CDRX",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "Final Exam Problem 3.5 (Page 29 - 12 Points)",
        "question": "Design a multi-range DC ammeter with ranges $I_1 = 1\\text{ mA}$, $I_2 = 50\\text{ mA}$, and $I_3 = 1\\text{ A}$ using a galvanometer with internal resistance $R_g = 100\\ \\Omega$, full-scale deflection current $I_g = 1\\ \\mu\\text{A}$, and damping resistor $\\text{CDRX} = 400\\ \\Omega$ in series with coil. Find shunt resistors $R_1, R_2, R_3$.\n" + make_img_html('exam_p29_ammeter_cdrx.png', 'Exam Problem 3.5: Multi-Range DC Ammeter with CDRX Circuit'),
        "formula": "$R_{\\text{branch}} = R_g + \\text{CDRX} = 500\\ \\Omega$\n$V_{\\text{branch}} = I_g \\times R_{\\text{branch}} = 1\\ \\mu\\text{A} \\times 500\\ \\Omega = 500\\ \\mu\\text{V} = 5 \\times 10^{-4}\\text{ V}$\n$R_s = \\frac{V_{\\text{branch}}}{I - I_g}$",
        "solution": "Total meter branch resistance: $R_{\\text{branch}} = 100 + 400 = 500\\ \\Omega$.\nBranch voltage at FSD: $V_m = 1\\times 10^{-6}\\text{ A} \\times 500\\ \\Omega = 5\\times 10^{-4}\\text{ V}$ ($0.5\\text{ mV}$).\n\n1. Range 1 ($I = 1\\text{ mA} = 10^{-3}\\text{ A}$):\n$$R_1 = \\frac{5 \\times 10^{-4}\\text{ V}}{10^{-3} - 10^{-6}} \\approx \\frac{5\\times 10^{-4}}{10^{-3}} = 0.5\\ \\Omega$$\n\n2. Range 2 ($I = 50\\text{ mA} = 50 \\times 10^{-3}\\text{ A}$):\n$$R_2 = \\frac{5 \\times 10^{-4}\\text{ V}}{50 \\times 10^{-3}} = 0.01\\ \\Omega$\n\n3. Range 3 ($I = 1\\text{ A}$):\n$$R_3 = \\frac{5 \\times 10^{-4}\\text{ V}}{1\\text{ A}} = 5 \\times 10^{-4}\\ \\Omega = 0.0005\\ \\Omega = 0.5\\text{ m}\\Omega$$.",
        "pitfall": "Exact 12-point exam problem! As measurement range increases, the required shunt resistance decreases proportionally.",
        "mastered": False
    },
    {
        "id": "CH10-06",
        "chapter": "Chapter 10",
        "topic": "Ayrton Shunt (Universal Shunt) for Multi-Range Ammeter",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 18 & Slide P.12",
        "question": "What is an Ayrton (Universal) Shunt in a multi-range DC ammeter? Why is it superior to individual switch shunts?\n" + make_img_html('multirange_ammeter.png', 'Individual Shunts vs Ayrton Universal Shunt Circuit'),
        "formula": "Series chain of resistors: $R_1 + R_2 + R_3$",
        "solution": "1. Structure: Employs a continuous series string of tapped resistors ($R_1, R_2, R_3$) permanently connected across the meter movement.\n2. Decisive Advantage over Individual Shunts: When switching between ranges, rotary switch contacts momentarily break. With individual shunts, the meter coil is momentarily left unshunted and exposed to the full line current, burning out the delicate hair spring. With an Ayrton shunt, the meter is NEVER left open or unshunted during range switching.",
        "pitfall": "Always mention the safety feature: the meter coil is never open-circuited while switching current ranges.",
        "mastered": False
    },
    {
        "id": "CH10-07",
        "chapter": "Chapter 10",
        "topic": "Multi-Range Voltmeter (Individual vs Tapped Series)",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 20 & Lecture Slides",
        "question": "Compare Individual Multipliers versus Series Tapped Multipliers for designing a multi-range DC voltmeter. State the resistance formulas for ranges $V_1, V_2, V_3$.\n" + make_img_html('multirange_voltmeter.png', 'Multi-Range Voltmeter: Individual vs Series Tapped Multipliers'),
        "formula": "Tapped: $R_1 = \\frac{V_1}{I_g} - R_g, R_2 = \\frac{V_2 - V_1}{I_g}, R_3 = \\frac{V_3 - V_2}{I_g}$",
        "solution": "1. Method 1: Individual Multipliers (Each range uses a separate dedicated resistor):\n$$R_{s1} = \\frac{V_1}{I_g} - R_g,\\quad R_{s2} = \\frac{V_2}{I_g} - R_g,\\quad R_{s3} = \\frac{V_3}{I_g} - R_g$$\n\n2. Method 2: Series Tapped Multipliers (Resistors connected in cumulative series chain):\n- Range 1 ($V_1$): $R_1 = \\frac{V_1}{I_g} - R_g$\n- Range 2 ($V_2$): $R_2 = \\frac{V_2 - V_1}{I_g}$\n- Range 3 ($V_3$): $R_3 = \\frac{V_3 - V_2}{I_g}$\nTotal series resistance for range 3 is $R_1 + R_2 + R_3$.\n\nBenefit of Tapped Series: Saves board space and requires standard lower resistor values.",
        "pitfall": "For tapped series, remember each resistor handles only the INCREMENTAL voltage $\\Delta V = V_k - V_{k-1}$, not the total voltage from zero.",
        "mastered": False
    },
    {
        "id": "CH10-08",
        "chapter": "Chapter 10",
        "topic": "Series-Type Ohmmeter Working Principle & Scale",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 17 & Lecture Slides",
        "question": "Explain the operating principle of a Series-Type Ohmmeter. Describe its zero-calibration, open-circuit state, and why its scale is reversed and non-linear.\n" + make_img_html('ohmmeter_circuits.png', 'Series-Type vs Shunt-Type Ohmmeter Circuits'),
        "formula": "$I_m = \\frac{E}{R_1 + R_2 + R_g + R_x}$\n$R_h = R_{\\text{in}} = R_1 + R_2 + R_g$",
        "solution": "1. Schematic: Battery $E$, current limiting resistor $R_1$, zero-adjust rheostat $R_2$, galvanometer $R_g$, and unknown resistor $R_x$ in series.\n2. Zero Calibration ($R_x = 0$, leads shorted):\nFull-scale deflection current flows: $I_{\\text{fsd}} = \\frac{E}{R_{\\text{in}}}$. Pointer deflects to FAR RIGHT ($0\\ \\Omega$).\n3. Open Circuit ($R_x = \\infty$, leads open):\nZero current flows: $I = 0$. Pointer stays at FAR LEFT ($\\infty\\ \\Omega$).\n4. Half-Scale Deflection ($I_m = 0.5 I_{\\text{fsd}}$):\nOccurs exactly when $R_x = R_{\\text{in}} = R_h$. Half-scale resistance equals the meter's internal circuit resistance.\n5. Non-linear Scale: Current is inversely proportional to $R_x$, causing severe crowding at the left (high resistance) end.",
        "pitfall": "Series ohmmeter scale runs BACKWARDS: $0\\ \\Omega$ is at the right, and $\\infty\\ \\Omega$ is at the left!",
        "mastered": False
    },
    {
        "id": "CH10-09",
        "chapter": "Chapter 10",
        "topic": "Shunt-Type Ohmmeter for Low Resistance Measurement",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 17 & Lecture Slides",
        "question": "Explain the operating principle of a Shunt-Type Ohmmeter. Why is its scale direction opposite to a series ohmmeter, and what resistance range is it suited for?\n" + make_img_html('ohmmeter_circuits.png', 'Shunt Ohmmeter Schematic'),
        "formula": "Scale: $0\\ \\Omega$ at LEFT, $\\infty\\ \\Omega$ at RIGHT",
        "solution": "1. Structure: The unknown resistor $R_x$ is connected in PARALLEL (shunt) across the galvanometer movement, with a switch $S$ to disconnect the internal battery when not measuring.\n2. Zero State ($R_x = 0$):\n$R_x$ completely shorts out the galvanometer coil. Meter current is zero ($I_m = 0$). Pointer stays at FAR LEFT ($0\\ \\Omega$).\n3. Open Circuit ($R_x = \\infty$):\nNo current is shunted away; full current flows through the meter coil ($I_m = I_{\\text{fsd}}$). Pointer deflects to FAR RIGHT ($\\infty\\ \\Omega$).\n4. Application: Specifically suited for measuring very LOW resistances ($0.5\\ \\Omega - 100\\ \\Omega$), such as motor windings and transformer coils.",
        "pitfall": "Remember: Series ohmmeter has $0\\ \\Omega$ on RIGHT; Shunt ohmmeter has $0\\ \\Omega$ on LEFT (same direction as voltmeter/ammeter).",
        "mastered": False
    },
    {
        "id": "CH10-10",
        "chapter": "Chapter 10",
        "topic": "Multi-Range Ohmmeter & Zero-Adjust Rheostat Function",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 17 & Lecture Slides",
        "question": "In a multi-range analog multimeter (ohmmeter), why is a Zero-Adjust Rheostat ($R_{\\text{adj}}$) necessary, and why must the user re-zero when switching ranges ($R\\times 1, R\\times 10, R\\times 1\\text{k}$)?",
        "formula": "Compensates for battery aging: $E \\downarrow$ and $r_{\\text{bat}} \\uparrow$",
        "solution": "1. Battery Aging Compensation: As internal chemical batteries are used and age, their open-circuit terminal voltage gradually decreases ($E \\downarrow$) and their internal resistance increases ($r_b \\uparrow$). Without adjustment, full-scale deflection cannot be reached when shorting leads.\n2. Zero-Adjust Rheostat Function: By turning the thumbwheel rheostat, the user varies internal resistance so that $I_{\\text{fsd}} = E / R_{\\text{in}}$ is precisely restored at $R_x = 0$.\n3. Range Switching: Different ranges ($R\\times 1, R\\times 10, R\\times 100$) switch in different shunt and series resistors, altering total internal circuit resistance $R_{\\text{in}}$. Therefore, the meter MUST be re-zeroed every time a new range is selected.",
        "pitfall": "Always short the test leads and adjust to $0\\ \\Omega$ before making any resistance measurement!",
        "mastered": False
    },
    {
        "id": "CH10-11",
        "chapter": "Chapter 10",
        "topic": "Rectifier AC Voltmeter & Form Factor",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 3.2 (Page 30 - 2 Points)",
        "question": "What is the operating principle of a Rectifier AC Voltmeter, and how does it relate average rectified DC to the indicated RMS value?",
        "formula": "$V_{\\text{rms}} = 1.11 \\times V_{\\text{avg}}$ (for pure sine waves)",
        "solution": "1. Principle: Uses semiconductor diodes (half-wave or full-wave bridge rectifier) to convert an alternating voltage into unidirectional pulsating DC, which is then measured by a standard PMMC movement.\n2. Scale Calibration: PMMC inherently responds to average value ($V_{\\text{avg}} = \\frac{2}{\\pi} V_m = 0.636 V_m$). The meter scale is multiplied by the Sine Wave Form Factor:\n$$\\text{Form Factor} = \\frac{V_{\\text{rms}}}{V_{\\text{avg}}} = \\frac{V_m / \\sqrt{2}}{2 V_m / \\pi} = \\frac{\\pi}{2\\sqrt{2}} = 1.11$$\n3. Reading Distortion: If the measured waveform is not a pure sine wave (e.g. square, triangle, or inverter PWM harmonics), the reading will have significant waveform error.",
        "pitfall": "A rectifier meter reads accurately ONLY for pure sine waves. For distorted waveforms, a True-RMS meter is required.",
        "mastered": False
    },
    {
        "id": "CH10-12",
        "chapter": "Chapter 10",
        "topic": "Wattmeter Connection to Load (Current vs Potential Coils)",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 3.4 (Page 30 - 4 Points)",
        "question": "In an electrodynamometer Wattmeter, how must the Current Coil and Potential (Voltage) Coil be connected into the circuit to measure load power?\n" + make_img_html('exam_p30_wattmeter.png', 'Exam Problem 3.4: Wattmeter Terminal Connections (1-2 and 3-5)'),
        "formula": "Current Coil (1-2) in SERIES | Potential Coil (3-5) in PARALLEL",
        "solution": "1. Current Coil (Terminals 1 - 2):\n- Made of thick copper wire with very low resistance.\n- Must be connected in SERIES with the load to carry full load current ($I_{\\text{load}}$).\n\n2. Potential / Voltage Coil (Terminals 3 - 5):\n- Made of many turns of fine wire in series with a high non-inductive resistor.\n- Must be connected in PARALLEL across the load to sense load voltage ($V_{\\text{load}}$).\n\n3. Deflection Torque:\n$$T_d \\propto V_{\\text{load}} I_{\\text{load}} \\cos\\theta = P\\quad [\\text{W}]$$.",
        "pitfall": "Never connect the current coil in parallel across the supply—it has virtually zero resistance and will cause an explosive short circuit!",
        "mastered": False
    },
    {
        "id": "CH10-13",
        "chapter": "Chapter 10",
        "topic": "Energy Meter (Kilowatt-Hour Meter) Function",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 3.3 (Page 30 - 2 Points)",
        "question": "What is an induction-type Energy Meter (Watt-hour meter) and what physical quantity and unit does it measure?",
        "formula": "$\\text{Energy} = \\int P(t) dt\\quad [\\text{kW}\\cdot\\text{h}]$",
        "solution": "1. Purpose: An integrating instrument that measures total electrical energy consumed over a period of time, commonly used by MEA/PEA for customer billing.\n2. Quantity & Unit: Measures active electrical energy in units of Kilowatt-hours ($\\text{kW}\\cdot\\text{h}$ or 'units').\n3. Principle: Eddy current induction motor principle driving a rotating aluminum disk geared to a register counter.",
        "pitfall": "A wattmeter measures instantaneous Power ($W$), while an energy meter measures cumulative Energy ($W \\cdot h$ or $kW \\cdot h$). Do not confuse the two!",
        "mastered": False
    },
    {
        "id": "CH10-14",
        "chapter": "Chapter 10",
        "topic": "Instrument Accuracy, Precision, and Error Formulas",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 15 & Slide P.3-5",
        "question": "Define and write the mathematical formulas for: (1) Absolute Error, (2) Relative Error (%), (3) Relative Accuracy, (4) Precision in electrical measurements.",
        "formula": "$\\epsilon = |Y_n - X_n|$\n$\\%\\text{Error} = \\frac{|Y_n - X_n|}{Y_n} \\times 100\\%$\n$\\%\\text{Accuracy} = 100\\% - \\%\\text{Error}$\n$\\text{Precision} = 1 - \\left|\\frac{X_n - \\bar{X}_n}{\\bar{X}_n}\\right|$",
        "solution": "1. Absolute Error ($\\epsilon$): Difference between true/expected value ($Y_n$) and measured value ($X_n$):\n$$\\epsilon = |Y_n - X_n|$$\n\n2. Percent Relative Error ($\\epsilon_r$):\n$$\\%\\text{Error} = \\frac{|Y_n - X_n|}{Y_n} \\times 100\\%$$\n\n3. Percent Accuracy ($A$):\n$$\\%\\text{Accuracy} = \\left(1 - \\frac{|Y_n - X_n|}{Y_n}\\right) \\times 100\\% = 100\\% - \\%\\text{Error}$$\n\n4. Precision: Measure of repeatability/consistency across multiple readings:\n$$\\text{Precision} = 1 - \\left|\\frac{X_n - \\bar{X}_n}{\\bar{X}_n}\\right|$$\nwhere $\\bar{X}_n$ is the arithmetic average of all repeated measurements.",
        "pitfall": "High precision does NOT guarantee high accuracy! A miscalibrated instrument can give highly repeatable (precise) readings that are completely wrong (inaccurate).",
        "mastered": False
    },

    # =========================================================================
    # CHAPTER 11 / SEMICONDUCTORS (12 CARDS)
    # =========================================================================
    {
        "id": "SEMI-01",
        "chapter": "Semiconductors",
        "topic": "Intrinsic vs Extrinsic, N-Type vs P-Type Semiconductors",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.a (Page 24 - 2 Points)",
        "question": "Explain the difference between Intrinsic and Extrinsic semiconductors. Compare N-type and P-type semiconductors in terms of dopants and majority carriers.",
        "formula": "N-type: Group V dopant (electrons) | P-type: Group III dopant (holes)",
        "solution": "1. Intrinsic Semiconductor: Chemically pure semiconductor (e.g. pure Silicon Si, Germanium Ge) with equal concentrations of free electrons and holes ($n = p = n_i$). Very low electrical conductivity at room temperature.\n2. Extrinsic Semiconductor: Doped with impurity atoms to dramatically enhance conductivity:\n- N-type: Pure Si doped with Pentavalent atoms (Group V: Phosphorus P, Arsenic As, Antimony Sb with 5 valence electrons). Majority carriers = Electrons; Minority carriers = Holes.\n- P-type: Pure Si doped with Trivalent atoms (Group III: Boron B, Gallium Ga, Indium In with 3 valence electrons). Majority carriers = Holes; Minority carriers = Electrons.",
        "pitfall": "Exact 2-point exam question from page 24! Memory aid: P in Pentavalent gives Negative electrons; T in Trivalent gives Positive holes.",
        "mastered": False
    },
    {
        "id": "SEMI-02",
        "chapter": "Semiconductors",
        "topic": "MOSFET Definition, Properties & Uses",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.b (Page 24 - 2 Points)",
        "question": "What is a MOSFET? State its essential electrical characteristics and its primary application/advantage in electronic circuits.",
        "formula": "Metal-Oxide-Semiconductor Field-Effect Transistor (Voltage-controlled)",
        "solution": "1. Definition: Metal-Oxide-Semiconductor Field-Effect Transistor, a unipolar transistor where current transport involves only majority carriers.\n2. Characteristics: Voltage-controlled device with extremely high input impedance ($10^{10} - 10^{15}\\ \\Omega$) due to the insulating $SiO_2$ gate oxide layer. Gate draws virtually zero DC current ($I_G \\approx 0$).\n3. Application / Benefit: Uses gate electric field to modulate and control current flow between Drain and Source. Widely used for high-speed switching, microprocessors (CMOS), and power converters.",
        "pitfall": "Exact question from page 24: State that it is a unipolar transistor whose output current is controlled by an electric field / voltage.",
        "mastered": False
    },
    {
        "id": "SEMI-03",
        "chapter": "Semiconductors",
        "topic": "Integrated Circuit (IC) Definition & Significance",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4.c (Page 24 - 1 Point)",
        "question": "What is an Integrated Circuit (IC)? What is its main function and significance in modern electronics?",
        "formula": "Microchip integrating transistors, diodes, resistors on single silicon die",
        "solution": "1. Definition: A miniaturized electronic circuit consisting of thousands to billions of interconnected semiconductor devices (transistors, diodes, resistors, capacitors) fabricated on a single tiny substrate of semiconductor material (silicon wafer).\n2. Significance: Enables massive reduction in physical size, weight, and power consumption, while offering dramatic improvements in reliability, operating speed, and mass-production cost.",
        "pitfall": "Exact 1-point question from past exam page 24.",
        "mastered": False
    },
    {
        "id": "SEMI-04",
        "chapter": "Semiconductors",
        "topic": "PN Junction Diode Biasing & Knee Voltages",
        "level": "Level 1 (Concept)",
        "exam_reference": "Lecture Slides Ch.11 & Jele Formula Sheet",
        "question": "Compare Forward Bias and Reverse Bias in a PN junction diode. State the barrier potential (knee voltage) values for Silicon (Si) and Germanium (Ge).\n" + make_img_html('diode_circuits.png', 'PN Diode Biasing, Forward Conduction and LED Circuit'),
        "formula": "Forward: P to +, N to - (conduction) | Reverse: P to -, N to + (blocking)\n$V_k(\\text{Si}) = 0.7\\text{ V}, V_k(\\text{Ge}) = 0.3\\text{ V}$",
        "solution": "1. Forward Bias: External positive connected to P-side (anode), negative to N-side (cathode). Overcomes internal barrier, depletion width shrinks, and heavy forward current flows once applied voltage exceeds knee voltage.\n2. Reverse Bias: Positive to N-side, negative to P-side. Depletion layer widens, blocking current except for tiny temperature-dependent reverse saturation leakage current ($I_0$).\n3. Knee Voltages:\n- Silicon (Si): $V_{\\text{knee}} = 0.7\\text{ V}$\n- Germanium (Ge): $V_{\\text{knee}} = 0.3\\text{ V}$.",
        "pitfall": "In almost all circuit analysis, Silicon diode forward voltage drop is taken as 0.7 V.",
        "mastered": False
    },
    {
        "id": "SEMI-05",
        "chapter": "Semiconductors",
        "topic": "Diode Shockley Equation, Dynamic Resistance & LED Circuit",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "FunE-final.pdf Page 22 & Lecture Slides",
        "question": "1. Write the Shockley diode equation and define thermal voltage $V_T$ at room temperature ($25^\\circ\\text{C}$). 2. For an LED circuit with $V_{CC} = 5\\text{ V}$, LED knee voltage $V_{AK} = 2.0\\text{ V}$, find the required current limiting resistor $R$ to maintain $I_D = 15\\text{ mA}$.",
        "formula": "$I_D = I_S(e^{V_D / n V_T} - 1)$\n$V_T = 25\\text{ mV}$\n$R = \\frac{V_{CC} - V_{AK}}{I_D}$",
        "solution": "1. Shockley Diode Equation:\n$$I_D = I_S\\left(e^{\\frac{V_D}{n V_T}} - 1\\right)$$\nwhere $I_S$ is reverse saturation current, $n$ is ideality factor ($1 \\le n \\le 2$), and thermal voltage $V_T = \\frac{k T}{q} \\approx 25\\text{ mV}$ (or $26\\text{ mV}$) at room temperature ($298\\text{ K}$).\n\n2. Dynamic / Small-Signal Resistance:\n$$r_d = \\frac{n V_T}{I_D} = \\frac{25\\text{ mV}}{I_D}$$\n\n3. LED Limiting Resistor Calculation:\n$$R = \\frac{V_{CC} - V_{AK}}{I_D} = \\frac{5.0 - 2.0\\text{ V}}{15 \\times 10^{-3}\\text{ A}} = \\frac{3.0}{0.015} = 200\\ \\Omega$$.",
        "pitfall": "Never connect an LED directly to a voltage supply without a series resistor; without current limiting, it will burn out instantaneously.",
        "mastered": False
    },
    {
        "id": "SEMI-06",
        "chapter": "Semiconductors",
        "topic": "Zener Diode Voltage Regulator Operation",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "Lecture Slides Ch.11 & Jele Formula Sheet",
        "question": "Explain how a Zener Diode regulates voltage. A $10\\text{ V}$ Zener diode is connected across a variable input voltage $V_{\\text{in}} = 15\\text{ V}$ with a series resistor $R_s = 200\\ \\Omega$ and load resistance $R_L = 1\\text{ k}\\Omega$. Calculate Zener current $I_Z$.",
        "formula": "$I_s = \\frac{V_{\\text{in}} - V_Z}{R_s}$\n$I_L = V_Z / R_L$\n$I_Z = I_s - I_L$",
        "solution": "1. Principle: Operates in the REVERSE BREAKDOWN region where voltage remains virtually constant ($V_Z$) across wide variations in reverse current.\n\n2. Calculation:\n- Regulated load voltage: $V_L = V_Z = 10\\text{ V}$\n- Source branch current: $I_s = \\frac{V_{\\text{in}} - V_Z}{R_s} = \\frac{15 - 10}{200} = \\frac{5}{200} = 25\\text{ mA}$\n- Load branch current: $I_L = \\frac{V_Z}{R_L} = \\frac{10\\text{ V}}{1,000\\ \\Omega} = 10\\text{ mA}$\n- Zener diode current: $I_Z = I_s - I_L = 25\\text{ mA} - 10\\text{ mA} = 15\\text{ mA}$.",
        "pitfall": "For regulation to hold, $I_Z$ must remain greater than knee current $I_{Z(min)}$ and less than maximum power rating current $I_{Z(max)} = P_Z / V_Z$.",
        "mastered": False
    },
    {
        "id": "SEMI-07",
        "chapter": "Semiconductors",
        "topic": "Solar Photovoltaic Cell Equivalent Circuit & Parameters",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 22 & Lecture Slides",
        "question": "Draw the equivalent circuit of a photovoltaic solar cell and define: Open-circuit voltage ($V_{oc}$), Short-circuit current ($I_{sc}$), Maximum power point ($P_{mp}$), and Fill Factor (FF).",
        "formula": "$\\text{FF} = \\frac{V_{mp} I_{mp}}{V_{oc} I_{sc}} = \\frac{P_{\\text{max}}}{V_{oc} I_{sc}}$",
        "solution": "1. Equivalent Circuit: Light-generated constant current source $I_L$ in parallel with a forward-biased PN junction diode $D$, parallel shunt leakage resistance $R_{sh}$, and series contact resistance $R_s$.\n\n2. Key Parameters:\n- $V_{oc}$: Open-circuit terminal voltage under illumination with zero load current (typically $0.6 - 0.7\\text{ V}$ per silicon cell).\n- $I_{sc}$: Short-circuit current when terminals are connected together ($V = 0$).\n- $P_{\\text{max}} = V_{mp} \\times I_{mp}$: Point on I-V curve where power output is maximized.\n- Fill Factor (FF): Ratio measuring the squareness of the I-V curve: $\\text{FF} = \\frac{V_{mp} I_{mp}}{V_{oc} I_{sc}}$, typically $0.75 - 0.85$ for high-quality silicon solar cells.",
        "pitfall": "A single silicon solar cell produces only ~0.6 V, so 60 to 72 cells are connected in series inside a panel to generate usable voltages (30-40 V).",
        "mastered": False
    },
    {
        "id": "SEMI-08",
        "chapter": "Semiconductors",
        "topic": "BJT Transistor Terminal Currents & Gain Relations",
        "level": "Level 1 (Concept)",
        "exam_reference": "Final Exam Problem 4 & Lecture Slides",
        "question": "State the fundamental KCL current relationship between Emitter ($I_e$), Collector ($I_c$), and Base ($I_b$) currents in a BJT, and the relationships with current gains $\\beta$ and $\\alpha$.\n" + make_img_html('transistor_bjt_mosfet.png', 'BJT NPN/PNP Symbols, Terminal Currents and MOSFET Structure'),
        "formula": "$I_e = I_c + I_b$\n$I_c = \\beta I_b$\n$\\alpha = \\frac{\\beta}{\\beta + 1} = \\frac{I_c}{I_e}$",
        "solution": "1. Terminal Current Law (KCL): Emitter current is the sum of collector and base currents:\n$$I_e = I_c + I_b$$\n2. Common-Emitter DC Current Gain ($\\beta$ or $h_{FE}$):\n$$I_c = \\beta I_b$$\nTypically $\\beta$ ranges from $50$ to $300$.\n3. Common-Base DC Current Gain ($\\alpha$):\n$$\\alpha = \\frac{I_c}{I_e} = \\frac{\\beta}{\\beta + 1}$$\nTypically $\\alpha \\approx 0.98 - 0.998$ (slightly less than 1.0).",
        "pitfall": "Because base current $I_b$ is tiny (typically $\\mu\\text{A}$), $I_e \\approx I_c$ in many rough approximations, but NEVER forget $I_e = I_c + I_b$ in exact KCL equations.",
        "mastered": False
    },
    {
        "id": "SEMI-09",
        "chapter": "Semiconductors",
        "topic": "BJT Transistor as an Electronic Switch",
        "level": "Level 1 (Concept)",
        "exam_reference": "FunE-final.pdf Page 23 & Lecture Slides",
        "question": "Explain how a BJT operates as an electronic switch. What conditions place the transistor into Cutoff (Switch OFF) versus Saturation (Switch ON)?",
        "formula": "Cutoff: $I_B = 0, V_{CE} = V_{CC}$ | Saturation: $I_B \\ge I_{C(\\text{sat})}/\\beta, V_{CE} \\approx 0.2\\text{ V}$",
        "solution": "1. Cutoff Region (Switch OPEN / OFF):\n- Condition: Base-Emitter voltage $V_{BE} < 0.7\\text{ V}$ (zero base drive).\n- Currents: $I_B = 0, I_C = 0$.\n- Output Voltage: $V_{CE} = V_{CC}$ (full supply voltage across open collector-emitter terminals).\n\n2. Saturation Region (Switch CLOSED / ON):\n- Condition: Base is heavily overdriven with $I_B \\ge \\frac{I_{C(\\text{sat})}}{\\beta}$, where $I_{C(\\text{sat})} \\approx \\frac{V_{CC}}{R_C}$.\n- Currents: Both junctions are forward-biased, collector current reaches maximum limited by load.\n- Output Voltage: $V_{CE} = V_{CE(\\text{sat})} \\approx 0.2\\text{ V}$ (acts like a closed switch with minimal drop).",
        "pitfall": "When using a transistor as a switch, always ensure $I_B > I_{C(sat)} / \\beta$ (overdrive factor of 2 to 3) to guarantee deep saturation.",
        "mastered": False
    },
    {
        "id": "SEMI-10",
        "chapter": "Semiconductors",
        "topic": "BJT vs FET / MOSFET Detailed Comparison",
        "level": "Level 1 (Concept)",
        "exam_reference": "Ch10IntroSemiDevices.pdf Slide P.21",
        "question": "Provide a comprehensive comparison between BJT (Bipolar Junction Transistor) and FET / MOSFET across 6 electrical parameters.",
        "formula": "BJT (Current-controlled, Bipolar) vs FET (Voltage-controlled, Unipolar)",
        "solution": "Direct exam comparison table from Slide 21:\n\n1. Control Mechanism:\n- BJT: Current-controlled device ($I_c = \\beta I_b$).\n- FET/MOSFET: Voltage-controlled device ($I_D = g_m V_{GS}$).\n\n2. Input Impedance:\n- BJT: Low ($1\\text{ k}\\Omega - 10\\text{ k}\\Omega$).\n- FET/MOSFET: Extremely high ($10^{10} - 10^{15}\\ \\Omega$, draws zero gate current).\n\n3. Charge Carriers:\n- BJT: Bipolar (conduction uses both electrons and holes).\n- FET: Unipolar (conduction uses only majority carriers in the channel).\n\n4. Noise:\n- BJT: More noisy.\n- FET: Less noisy.\n\n5. Thermal Stability:\n- BJT: Less temperature stable (prone to thermal runaway).\n- FET: More temperature stable.\n\n6. Physical Size on Chip:\n- BJT: Larger size.\n- MOSFET: Much smaller size (ideal for billion-transistor IC microchips).",
        "pitfall": "Memorize this exact comparison table from Slide 21! It appears regularly in theory short-answer and multiple-choice questions.",
        "mastered": False
    },
    {
        "id": "SEMI-11",
        "chapter": "Semiconductors",
        "topic": "Half-Wave vs Full-Wave Rectifiers & PIV Rating",
        "level": "Level 2 (Past Exam / Calc)",
        "exam_reference": "X-Class Review Page 7-9 & Lecture Slides",
        "question": "Compare Half-Wave Rectifier, Full-Wave Center-Tapped Rectifier, and Full-Wave Bridge Rectifier in terms of: (1) Number of diodes, (2) DC average output voltage ($V_{\\text{dc}}$), (3) Peak Inverse Voltage (PIV) rating across each diode.",
        "formula": "Half-wave: $V_{\\text{dc}} = \\frac{V_m}{\\pi} = 0.318 V_m$, $\\text{PIV} = V_m$\nFull-wave: $V_{\\text{dc}} = \\frac{2V_m}{\\pi} = 0.636 V_m$\nCenter-Tapped PIV = $2 V_m$ | Bridge PIV = $V_m$",
        "solution": "1. Half-Wave Rectifier:\n- 1 diode\n- $V_{\\text{dc}} = V_m / \\pi = 0.318 V_m$\n- Ripple frequency = $f_{\\text{in}}$ ($50\\text{ Hz}$)\n- $\\text{PIV} = V_m$.\n\n2. Full-Wave Center-Tapped Rectifier:\n- 2 diodes\n- $V_{\\text{dc}} = 2V_m / \\pi = 0.636 V_m$\n- Ripple frequency = $2 f_{\\text{in}}$ ($100\\text{ Hz}$)\n- $\\text{PIV} = 2 V_m$ (Diodes must withstand twice peak voltage!).\n\n3. Full-Wave Bridge Rectifier:\n- 4 diodes\n- $V_{\\text{dc}} = 2V_m / \\pi = 0.636 V_m$\n- Ripple frequency = $2 f_{\\text{in}}$ ($100\\text{ Hz}$)\n- $\\text{PIV} = V_m$ (Does NOT require expensive center-tapped transformer).",
        "pitfall": "Exam multiple choice favorite: Full-wave center-tapped diode requires $\\text{PIV} = 2V_m$, whereas Bridge diode requires only $\\text{PIV} = V_m$.",
        "mastered": False
    },
    {
        "id": "SEMI-12",
        "chapter": "Semiconductors",
        "topic": "Operational Amplifier (Op-Amp 741) Characteristics",
        "level": "Level 1 (Concept)",
        "exam_reference": "Ch10IntroSemiDevices.pdf Slide P.23-24",
        "question": "What is an Operational Amplifier (Op-Amp)? Identify its 5 primary terminal pins and state the 3 ideal Op-Amp properties.",
        "formula": "$V_{\\text{out}} = A_{OL}(V_+ - V_-)$",
        "solution": "1. Definition: A high-gain DC-coupled differential voltage amplifier integrated circuit (e.g. standard $\\mu\\text{A}741$).\n2. Terminal Pins:\n- Pin 2: Inverting input ($V_-$)\n- Pin 3: Non-inverting input ($V_+$)\n- Pin 6: Output ($V_{\\text{out}}$)\n- Pin 7: Positive power supply ($V_S^+$)\n- Pin 4: Negative power supply ($V_S^-$).\n\n3. Ideal Op-Amp Properties:\n- Infinite open-loop voltage gain: $A_{OL} = \\infty$\n- Infinite input impedance: $R_{\\text{in}} = \\infty$ (no input current, $I_+ = I_- = 0$)\n- Zero output impedance: $R_{\\text{out}} = 0$\n- Infinite bandwidth and zero input offset voltage.",
        "pitfall": "Virtual ground concept: When negative feedback is present, the Op-Amp drives the differential input to zero ($V_+ = V_-$).",
        "mastered": False
    }
]

print(f"Total flashcards crafted: {len(flashcards)}")

# 1. Generate Excel file
excel_path = r'c:\Users\thana\Desktop\Private-Flashcard\FunE_Master_Exam_Flashcards_Ch6_10.xlsx'
df = pd.DataFrame(flashcards)
df_export = df[['id', 'chapter', 'topic', 'level', 'exam_reference', 'question', 'formula', 'solution', 'pitfall', 'mastered']]
df_export.columns = [
    'Card ID', 'Chapter', 'Topic', 'Level / Difficulty', 'Exam Reference', 
    'Question / Problem', 'Formula / Key Concept', 'Step-by-Step Solution', 'Pitfall / Pro Tip', 'Mastered'
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Flashcards"

headers = list(df_export.columns)
ws.append(headers)

for row in df_export.itertuples(index=False):
    ws.append(list(row))

header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

ws.row_dimensions[1].height = 28
for col_num, cell in enumerate(ws[1], 1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row_num in range(2, len(df_export) + 2):
    ws.row_dimensions[row_num].height = 70
    is_even = (row_num % 2 == 0)
    row_fill = PatternFill(start_color="F8FAFC" if is_even else "FFFFFF", fill_type="solid")
    
    for col_num, cell in enumerate(ws[row_num], 1):
        cell.fill = row_fill
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)
        if col_num in [1, 2, 4, 10]:
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

col_widths = {
    'A': 12, 'B': 18, 'C': 28, 'D': 22, 'E': 28, 
    'F': 45, 'G': 32, 'H': 50, 'I': 34, 'J': 12
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# 1. Generate JSON and JS files first
json_path = r'c:\Users\thana\Desktop\Private-Flashcard\data\flashcards.json'
js_path = r'c:\Users\thana\Desktop\Private-Flashcard\data\flashcards.js'

def format_card_answer_html(formula_text, solution_text, pitfall_text):
    # 1. Formula Box: Every formula line on its own row with clear vertical padding
    f_lines = [l.strip() for l in formula_text.split('\n') if l.strip()]
    f_rows = "".join(f"<div class='formula-row py-1'>{line}</div>" for line in f_lines)
    
    # 2. Solution Box: Every step, bullet, and equation on its own row
    s_lines = [l.strip() for l in solution_text.split('\n') if l.strip()]
    s_rows = []
    for line in s_lines:
        if line.startswith('$$') and line.endswith('$$'):
            s_rows.append(f"<div class='solution-eq my-2 py-1 text-center font-medium'>{line}</div>")
        elif any(line.startswith(f"{i}.") for i in range(1, 10)) or any(line.startswith(f"({ch})") for ch in 'abcd'):
            s_rows.append(f"<div class='solution-step font-bold text-slate-900 dark:text-slate-100 mt-2 mb-1'>{line}</div>")
        elif line.startswith('- '):
            s_rows.append(f"<div class='solution-bullet my-1 pl-3 text-slate-700 dark:text-slate-300'>{line}</div>")
        else:
            s_rows.append(f"<div class='solution-line my-1 text-slate-800 dark:text-slate-200'>{line}</div>")
    s_content = "".join(s_rows)
    
    # 3. Pitfall Box: Clean alert box
    return (
        f"<div class='space-y-3'>"
        f"<div class='formula-box bg-blue-50/90 dark:bg-blue-950/40 border-l-4 border-blue-600 px-3.5 py-2.5 rounded-r-lg shadow-xs'>"
        f"<div class='text-xs font-bold uppercase tracking-wider text-blue-700 dark:text-blue-300 mb-1.5'>Key Formula / Operating Principle:</div>"
        f"<div class='text-sm leading-relaxed'>{f_rows}</div>"
        f"</div>"
        f"<div class='solution-box bg-slate-50 dark:bg-slate-800/70 border-l-4 border-emerald-600 px-3.5 py-2.5 rounded-r-lg shadow-xs'>"
        f"<div class='text-xs font-bold uppercase tracking-wider text-emerald-700 dark:text-emerald-300 mb-1.5'>Step-by-Step Solution & Calculation:</div>"
        f"<div class='text-sm leading-relaxed'>{s_content}</div>"
        f"</div>"
        f"<div class='pitfall-box bg-amber-50/90 dark:bg-amber-950/40 border-l-4 border-amber-500 px-3.5 py-2 rounded-r-lg text-xs leading-relaxed text-amber-900 dark:text-amber-200'>"
        f"<span class='font-bold uppercase tracking-wider text-amber-700 dark:text-amber-400 mr-1'>⚠️ Exam Pitfall / Pro Tip:</span>"
        f"<span class='italic'>{pitfall_text}</span>"
        f"</div>"
        f"</div>"
    )

web_cards = []
for c in flashcards:
    web_cards.append({
        "id": c["id"],
        "chapter": c["chapter"],
        "topic": c["topic"],
        "difficulty": c["level"],
        "exam_reference": c["exam_reference"],
        "question": c["question"],
        "hint": f"Key Formula: {c['formula'].splitlines()[0]}" if c["formula"] else "Click to flip card for answer",
        "answer": format_card_answer_html(c['formula'], c['solution'], c['pitfall'])
    })

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(web_cards, f, ensure_ascii=False, indent=2)

with open(js_path, 'w', encoding='utf-8') as f:
    f.write("const FLASHCARDS_DATA = " + json.dumps(web_cards, ensure_ascii=False, indent=2) + ";\n")

print(f"Exported JSON ({len(web_cards)} cards) to: {json_path}")
print(f"Exported JS ({len(web_cards)} cards) to: {js_path}")

# 2. Save Excel file
ws.freeze_panes = "A2"
try:
    wb.save(excel_path)
    print(f"Exported Excel to: {excel_path}")
except PermissionError:
    alt_excel = r'c:\Users\thana\Desktop\Private-Flashcard\FunE_Master_Exam_Flashcards_Ch6_10_Master.xlsx'
    wb.save(alt_excel)
    print(f"Original Excel file was locked by another process. Saved to: {alt_excel}")